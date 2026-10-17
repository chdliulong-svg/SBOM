# blueprints/materials.py
"""
Global Master Materials Library Management, Structures, Auditing & Excel Batch Importer.
"""

import os
import io
import time
import json
from flask import Blueprint, request, jsonify, send_file
from werkzeug.utils import secure_filename
from services.database import get_db, get_material_tree
from services.backup_manager import create_backup_snapshot
from services.data_audit import get_duplicate_materials_report, check_single_material_duplicate
from services.master_importer import analyze_master_excel, execute_master_import
from services.logging_setup import get_logger
from blueprints.common import UPLOAD_DIR

logger = get_logger('sbom.materials')
materials_bp = Blueprint('materials', __name__)


# ----------------- Materials Master Query & Stats -----------------
@materials_bp.route('/api/materials', methods=['GET'])
def get_materials():
    search = request.args.get('q', '').strip()
    filter_type = request.args.get('filter_type', 'all').strip()
    conn = get_db()
    cursor = conn.cursor()

    sql = '''
    WITH child_counts AS (
        SELECT parent_material_no, COUNT(*) as cnt
        FROM material_structures
        GROUP BY parent_material_no
    ),
    usage_stats AS (
        SELECT b.material_no,
               COUNT(DISTINCT sub.project_id) as proj_cnt,
               COUNT(*) as total_cnt
        FROM bom_items b
        JOIN subsystems sub ON b.subsystem_id = sub.id
        GROUP BY b.material_no
    ),
    proj_refs AS (
        SELECT material_no,
               json_group_array(json_object('id', id, 'name', name, 'code', code)) as projects_json
        FROM (
            SELECT DISTINCT b.material_no, p.id, COALESCE(NULLIF(p.system_name, ''), NULLIF(p.code, ''), p.name) as name, p.code
            FROM bom_items b
            JOIN subsystems sub ON b.subsystem_id = sub.id
            JOIN projects p ON sub.project_id = p.id
            ORDER BY COALESCE(NULLIF(p.system_name, ''), p.code)
        )
        GROUP BY material_no
    )
    SELECT m.*,
           COALESCE(c.cnt, 0) as child_count,
           COALESCE(u.proj_cnt, 0) as usage_project_count,
           COALESCE(u.total_cnt, 0) as usage_total_count,
           pr.projects_json
    FROM materials m
    LEFT JOIN child_counts c ON c.parent_material_no = m.material_no
    LEFT JOIN usage_stats u ON u.material_no = m.material_no
    LEFT JOIN proj_refs pr ON pr.material_no = m.material_no
    WHERE 1=1
    '''
    params = []

    if search:
        query = f"%{search}%"
        sql += """ AND (
            m.material_no LIKE ? OR m.name LIKE ? OR m.description LIKE ?
            OR m.spec_model LIKE ? OR m.manufacturer LIKE ? OR m.drawing_no LIKE ? OR m.notes LIKE ?
            OR EXISTS (
                SELECT 1 FROM bom_items b
                JOIN subsystems sub ON b.subsystem_id = sub.id
                JOIN projects p ON sub.project_id = p.id
                WHERE b.material_no = m.material_no AND (p.name LIKE ? OR p.code LIKE ? OR p.system_name LIKE ?)
            )
        )"""
        params.extend([query, query, query, query, query, query, query, query, query, query])

    if filter_type == 'has_children':
        sql += " AND COALESCE(c.cnt, 0) > 0"
    elif filter_type == 'has_image':
        sql += " AND m.image_path IS NOT NULL AND m.image_path != ''"
    elif filter_type == 'has_manual':
        sql += " AND m.manual_path IS NOT NULL AND m.manual_path != ''"
    elif filter_type == 'referenced':
        sql += " AND COALESCE(u.total_cnt, 0) > 0"
    elif filter_type == 'unreferenced':
        sql += " AND COALESCE(u.total_cnt, 0) = 0"

    sql += " ORDER BY m.material_no"

    limit = request.args.get('limit', type=int)
    offset = request.args.get('offset', type=int)
    if limit is not None and limit > 0:
        sql += " LIMIT ?"
        params.append(limit)
        if offset is not None and offset >= 0:
            sql += " OFFSET ?"
            params.append(offset)
    cursor.execute(sql, tuple(params))
    raw_materials = cursor.fetchall()
    materials = []
    for row in raw_materials:
        d = dict(row)
        pj = d.get('projects_json')
        if pj:
            try:
                d['referenced_projects'] = json.loads(pj)
            except Exception:
                d['referenced_projects'] = []
        else:
            d['referenced_projects'] = []
        materials.append(d)
    conn.close()
    return jsonify(materials)


@materials_bp.route('/api/materials/stats', methods=['GET'])
def get_materials_stats():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT COUNT(*) FROM materials')
    total = cursor.fetchone()[0]
    cursor.execute('SELECT COUNT(DISTINCT parent_material_no) FROM material_structures')
    with_children = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM materials WHERE image_path IS NOT NULL AND image_path != ''")
    with_image = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM materials WHERE manual_path IS NOT NULL AND manual_path != ''")
    with_manual = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(DISTINCT material_no) FROM bom_items WHERE material_no IS NOT NULL AND material_no != ''")
    referenced_count = cursor.fetchone()[0]
    conn.close()
    return jsonify({
        'total': total,
        'total_materials': total,
        'with_children': with_children,
        'assemblies_count': with_children,
        'with_image': with_image,
        'images_count': with_image,
        'with_manual': with_manual,
        'manuals_count': with_manual,
        'referenced_count': referenced_count,
        'unreferenced_count': max(0, total - referenced_count)
    })


@materials_bp.route('/api/material/<material_no>/children', methods=['GET'])
def get_material_children_api(material_no):
    mat_no = material_no.strip()
    conn = get_db()
    cursor = conn.cursor()
    children = get_material_tree(cursor, mat_no)
    conn.close()
    return jsonify({
        'has_children': len(children) > 0,
        'count': len(children),
        'material_no': mat_no,
        'children': children
    })


@materials_bp.route('/api/material/save', methods=['POST'])
def save_material():
    data = request.json or {}
    mat_no = data.get('material_no', '').strip()
    old_mat_no = (data.get('old_material_no') or '').strip()
    if not mat_no:
        return jsonify({'error': '物料号不能为空'}), 400

    name = data.get('name', '').strip()
    description = data.get('description', '').strip()
    spec_model = data.get('spec_model', '').strip()
    manufacturer = data.get('manufacturer', '').strip()
    unit = data.get('unit', 'PC').strip() or 'PC'
    substitute = data.get('substitute', '').strip()
    drawing_no = data.get('drawing_no', '/').strip() or '/'
    notes = data.get('notes', '').strip()
    is_create = data.get('is_create', False)
    sync_bom = data.get('sync_bom', False)
    remove_image = bool(data.get('remove_image'))
    remove_manual = bool(data.get('remove_manual'))

    conn = get_db()
    cursor = conn.cursor()

    if is_create:
        cursor.execute('SELECT COUNT(*) FROM materials WHERE material_no = ?', (mat_no,))
        if cursor.fetchone()[0] > 0:
            conn.close()
            return jsonify({'error': f'物料号 [{mat_no}] 已存在于主库中，请勿重复创建！'}), 400

        cursor.execute('''
        INSERT INTO materials (material_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute, image_path, manual_path)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (mat_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute,
              data.get('image_path', '').strip(), data.get('manual_path', '').strip()))
    else:
        target_old = old_mat_no if old_mat_no else mat_no
        if target_old != mat_no:
            cursor.execute('SELECT COUNT(*) FROM materials WHERE material_no = ?', (mat_no,))
            if cursor.fetchone()[0] > 0:
                return jsonify({'error': f'新物料号 [{mat_no}] 已存在，无法将当前物料修改为此编号！'}), 400

            # 开启延迟外键检查，允许在同一事务内级联同步主表与引用子表（material_structures, bom_items）
            cursor.execute('PRAGMA defer_foreign_keys = ON;')

            cursor.execute('''
            UPDATE materials
            SET material_no = ?, name = ?, description = ?, spec_model = ?, manufacturer = ?, unit = ?, drawing_no = ?, notes = ?, substitute = ?
            WHERE material_no = ?
            ''', (mat_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute, target_old))

            cursor.execute('UPDATE material_structures SET parent_material_no = ? WHERE parent_material_no = ?', (mat_no, target_old))
            cursor.execute('UPDATE material_structures SET child_material_no = ? WHERE child_material_no = ?', (mat_no, target_old))
            cursor.execute('UPDATE bom_items SET material_no = ? WHERE material_no = ?', (mat_no, target_old))
            cursor.execute('UPDATE bom_items SET parent_material_no = ? WHERE parent_material_no = ?', (mat_no, target_old))
        else:
            cursor.execute('''
            INSERT INTO materials (material_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute, image_path, manual_path)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(material_no) DO UPDATE SET
                name=excluded.name,
                description=excluded.description,
                spec_model=excluded.spec_model,
                manufacturer=excluded.manufacturer,
                unit=excluded.unit,
                drawing_no=excluded.drawing_no,
                notes=excluded.notes,
                substitute=excluded.substitute,
                image_path=CASE WHEN excluded.image_path != '' THEN excluded.image_path ELSE materials.image_path END,
                manual_path=CASE WHEN excluded.manual_path != '' THEN excluded.manual_path ELSE materials.manual_path END
            ''', (mat_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute,
                  data.get('image_path', '').strip(), data.get('manual_path', '').strip()))

    if remove_image:
        cursor.execute("UPDATE materials SET image_path = '' WHERE material_no = ?", (mat_no,))
    if remove_manual:
        cursor.execute("UPDATE materials SET manual_path = '' WHERE material_no = ?", (mat_no,))

    synced_rows_count = 0
    if sync_bom:
        cursor.execute('''
        UPDATE bom_items
        SET unit = ?, notes = ?, substitute = ?, drawing_no = CASE WHEN ? != '/' THEN ? ELSE drawing_no END
        WHERE material_no = ?
          AND subsystem_id IN (
              SELECT s.id FROM subsystems s
              JOIN projects p ON s.project_id = p.id
              WHERE p.is_locked = 0
          )
        ''', (unit, notes, substitute, drawing_no, drawing_no, mat_no))
        synced_rows_count = cursor.rowcount

    conn.commit()
    return jsonify({'success': True, 'material_no': mat_no, 'synced_rows_count': synced_rows_count})


@materials_bp.route('/api/material/delete', methods=['POST'])
def delete_material():
    data = request.json or {}
    mat_no = data.get('material_no', '').strip()
    if not mat_no:
        return jsonify({'error': '未提供物料号'}), 400

    conn = get_db()
    cursor = conn.cursor()

    # 1. 检查外键引用：若该物料已被工程项目 BOM 引用，阻止直接删除，避免 IntegrityError 崩溃
    cursor.execute('''
        SELECT COUNT(DISTINCT p.id), GROUP_CONCAT(DISTINCT COALESCE(NULLIF(p.code, ''), p.name))
        FROM bom_items b
        JOIN subsystems s ON b.subsystem_id = s.id
        JOIN projects p ON s.project_id = p.id
        WHERE b.material_no = ?
    ''', (mat_no,))
    ref_row = cursor.fetchone()
    ref_count = ref_row[0] or 0
    if ref_count > 0:
        proj_names = ref_row[1] or '相关工程'
        return jsonify({
            'error': f'物料 [{mat_no}] 已被 {ref_count} 个工程项目引用（涉及工程：{proj_names}），无法从主库中直接删除！请先在相关工程中解除对该物料的引用。'
        }), 400

    try:
        create_backup_snapshot(reason='delete_material', note=f'删除通用物料 [{mat_no}] 前自动备份')
    except Exception as e_snap:
        logger.warning(f"Auto-backup before delete material warning: {e_snap}")

    cursor.execute('DELETE FROM material_structures WHERE parent_material_no = ? OR child_material_no = ?', (mat_no, mat_no))
    cursor.execute('DELETE FROM materials WHERE material_no = ?', (mat_no,))
    conn.commit()
    return jsonify({'success': True, 'message': f'物料 [{mat_no}] 已成功从全局通用物料主库中移除！'})


@materials_bp.route('/api/materials/batch-delete', methods=['POST'])
def batch_delete_materials():
    data = request.json or {}
    mat_nos = data.get('material_nos', [])
    if not isinstance(mat_nos, list) or not mat_nos:
        return jsonify({'error': '未提供待删除的物料号列表'}), 400

    clean_nos = [str(x).strip() for x in mat_nos if str(x).strip()]
    if not clean_nos:
        return jsonify({'error': '待删除物料号列表为空'}), 400

    conn = get_db()
    cursor = conn.cursor()
    placeholders = ','.join(['?'] * len(clean_nos))

    # 1. 检查外键引用：拦截已被工程 BOM 引用的物料
    cursor.execute(f'''
        SELECT b.material_no, COUNT(DISTINCT p.id), GROUP_CONCAT(DISTINCT COALESCE(NULLIF(p.code, ''), p.name))
        FROM bom_items b
        JOIN subsystems s ON b.subsystem_id = s.id
        JOIN projects p ON s.project_id = p.id
        WHERE b.material_no IN ({placeholders})
        GROUP BY b.material_no
    ''', clean_nos)
    in_use_rows = cursor.fetchall()
    if in_use_rows:
        in_use_details = [f"[{r[0]}]（工程: {r[2] or '未知'}）" for r in in_use_rows]
        detail_str = '、'.join(in_use_details[:5])
        if len(in_use_details) > 5:
            detail_str += f' 等共 {len(in_use_details)} 项'
        return jsonify({
            'error': f'选中的物料已被工程项目引用，禁止删除以保护数据一致性：\n{detail_str}。\n请先在对应工程中解除引用后再操作。'
        }), 400

    try:
        create_backup_snapshot(reason='batch_delete_materials', note=f'批量删除 {len(clean_nos)} 个通用物料前自动备份')
    except Exception as e_snap:
        logger.warning(f"Auto-backup before batch delete materials warning: {e_snap}")

    cursor.execute(f'DELETE FROM material_structures WHERE parent_material_no IN ({placeholders}) OR child_material_no IN ({placeholders})', clean_nos + clean_nos)
    cursor.execute(f'DELETE FROM materials WHERE material_no IN ({placeholders})', clean_nos)
    deleted_count = cursor.rowcount
    conn.commit()
    return jsonify({
        'success': True,
        'deleted_count': deleted_count,
        'message': f'成功从物料主库中批量删除 {len(clean_nos)} 项物料！'
    })


@materials_bp.route('/api/materials/clear', methods=['POST'])
def clear_materials():
    return jsonify({
        'success': False,
        'error': '清空物料主库功能已被系统彻底禁用并移除，以防止不可逆数据丢失！'
    }), 403


@materials_bp.route('/api/material/<path:material_no>/usage', methods=['GET'])
def get_material_usage(material_no):
    mat_no = material_no.strip()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    SELECT b.id, b.row_num, b.seq_num, b.qty, b.unit, b.level_num, b.level_str,
           s.id as subsystem_id, s.name as subsystem_name,
           p.id as project_id,
           COALESCE(NULLIF(p.system_name, ''), NULLIF(p.code, ''), p.name) as project_name,
           p.code as project_code,
           p.name as product_name
    FROM bom_items b
    JOIN subsystems s ON b.subsystem_id = s.id
    JOIN projects p ON s.project_id = p.id
    WHERE b.material_no = ?
    ORDER BY COALESCE(NULLIF(p.system_name, ''), p.name), s.name, b.row_num
    ''', (mat_no,))
    usage_list = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify({
        'material_no': mat_no,
        'count': len(usage_list),
        'usage': usage_list,
        'usages': usage_list
    })


@materials_bp.route('/api/material/<path:material_no>/structure', methods=['GET'])
def get_material_structure(material_no):
    mat_no = material_no.strip()
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    SELECT s.id, s.parent_material_no, s.child_material_no, s.child_level_num,
           s.child_level_str, s.sort_order, s.qty, s.unit, s.drawing_no, s.notes, s.substitute,
           s.child_level_num as level_num,
           COALESCE(m.name, '') as name,
           COALESCE(m.description, '') as description,
           COALESCE(m.spec_model, '') as spec_model,
           COALESCE(m.manufacturer, '') as manufacturer,
           COALESCE(m.name, '') as m_name,
           COALESCE(m.description, '') as m_desc,
           COALESCE(m.spec_model, '') as m_model,
           COALESCE(m.manufacturer, '') as m_mfr,
           m.unit as m_unit, m.image_path, m.manual_path
    FROM material_structures s
    LEFT JOIN materials m ON s.child_material_no = m.material_no
    WHERE s.parent_material_no = ?
    ORDER BY s.sort_order, s.id
    ''', (mat_no,))
    children = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify({
        'parent_material_no': mat_no,
        'count': len(children),
        'children': children
    })


@materials_bp.route('/api/material/structure/save', methods=['POST'])
def save_material_structure():
    data = request.json or {}
    struct_id = data.get('id')
    parent_mat = data.get('parent_material_no', '').strip()
    child_mat = data.get('child_material_no', '').strip()

    if not parent_mat or not child_mat:
        return jsonify({'error': '父物料号和子物料号均不能为空'}), 400

    if parent_mat.lower() == child_mat.lower():
        return jsonify({'error': '父物料号与子物料号不能相同'}), 400

    conn = get_db()
    cursor = conn.cursor()

    parent_lvl = 1
    cursor.execute('SELECT child_level_num FROM material_structures WHERE child_material_no = ?', (parent_mat,))
    p_struct = cursor.fetchone()
    if p_struct and p_struct[0]:
        parent_lvl = int(p_struct[0])
    else:
        cursor.execute('SELECT level_num FROM bom_items WHERE material_no = ? ORDER BY id DESC', (parent_mat,))
        b_row = cursor.fetchone()
        if b_row and b_row[0]:
            parent_lvl = int(b_row[0])

    lvl = parent_lvl + 1
    lvl_str = '•' * lvl
    qty = str(data.get('qty', '1')).strip()
    unit = data.get('unit', 'PC').strip() or 'PC'
    dwg = data.get('drawing_no', '/').strip()
    notes = data.get('notes', '').strip()
    substitute = data.get('substitute', '').strip()

    cursor.execute('SELECT COUNT(*) FROM materials WHERE material_no = ?', (child_mat,))
    if cursor.fetchone()[0] == 0:
        cursor.execute('''
        INSERT INTO materials (material_no, name, description, spec_model, manufacturer, unit)
        VALUES (?, ?, ?, ?, ?, ?)
        ''', (child_mat, data.get('m_name', '').strip(), data.get('m_desc', '').strip(),
              data.get('m_model', '').strip(), data.get('m_mfr', '').strip(), unit))

    if struct_id:
        cursor.execute('''
        UPDATE material_structures
        SET child_level_num = ?, child_level_str = ?, qty = ?, unit = ?, drawing_no = ?, notes = ?, substitute = ?
        WHERE id = ?
        ''', (lvl, lvl_str, qty, unit, dwg, notes, substitute, struct_id))
    else:
        cursor.execute('SELECT id FROM material_structures WHERE parent_material_no = ? AND child_material_no = ?',
                       (parent_mat, child_mat))
        existing = cursor.fetchone()
        if existing:
            cursor.execute('''
            UPDATE material_structures
            SET child_level_num = ?, child_level_str = ?, qty = ?, unit = ?, drawing_no = ?, notes = ?, substitute = ?
            WHERE id = ?
            ''', (lvl, lvl_str, qty, unit, dwg, notes, substitute, existing[0]))
        else:
            cursor.execute('SELECT COALESCE(MAX(sort_order), 0) + 1 FROM material_structures WHERE parent_material_no = ?', (parent_mat,))
            next_sort = cursor.fetchone()[0]
            cursor.execute('''
            INSERT INTO material_structures (parent_material_no, child_material_no, child_level_num, child_level_str, sort_order, qty, unit, drawing_no, notes, substitute)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (parent_mat, child_mat, lvl, lvl_str, next_sort, qty, unit, dwg, notes, substitute))

    conn.commit()
    conn.close()
    return jsonify({'success': True})


@materials_bp.route('/api/material/structure/delete', methods=['POST'])
def delete_material_structure():
    data = request.json or {}
    struct_id = data.get('id')
    if not struct_id:
        return jsonify({'error': '未提供结构ID'}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM material_structures WHERE id = ?', (struct_id,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@materials_bp.route('/api/material/structure/clear', methods=['POST'])
def clear_material_structures():
    data = request.json or {}
    parent_mat = data.get('parent_material_no', '').strip()
    if not parent_mat:
        return jsonify({'error': '未提供父物料号'}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM material_structures WHERE parent_material_no = ?', (parent_mat,))
    deleted_count = cursor.rowcount
    conn.commit()
    conn.close()
    return jsonify({
        'success': True,
        'deleted_count': deleted_count,
        'message': f"已成功清空总成 [{parent_mat}] 的全部下属子零部件（共 {deleted_count} 项）"
    })


@materials_bp.route('/api/material/structure/move', methods=['POST'])
def move_material_structure():
    data = request.json or {}
    struct_id = data.get('id')
    direction = data.get('direction', 'up')
    if not struct_id:
        return jsonify({'error': '未提供结构ID'}), 400

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT id, parent_material_no, sort_order FROM material_structures WHERE id = ?', (struct_id,))
    curr = cursor.fetchone()
    if not curr:
        conn.close()
        return jsonify({'error': '结构记录未找到'}), 404

    parent_mat = curr['parent_material_no']
    curr_order = curr['sort_order']

    cursor.execute('SELECT id, sort_order FROM material_structures WHERE parent_material_no = ? ORDER BY sort_order, id', (parent_mat,))
    all_children = cursor.fetchall()
    all_ids = [r['id'] for r in all_children]
    idx = all_ids.index(struct_id)

    if direction == 'up' and idx > 0:
        other_id = all_ids[idx - 1]
    elif direction == 'down' and idx < len(all_ids) - 1:
        other_id = all_ids[idx + 1]
    else:
        conn.close()
        return jsonify({'success': True})

    cursor.execute('SELECT sort_order FROM material_structures WHERE id = ?', (other_id,))
    other_order = cursor.fetchone()['sort_order']

    cursor.execute('UPDATE material_structures SET sort_order = ? WHERE id = ?', (other_order, struct_id))
    cursor.execute('UPDATE material_structures SET sort_order = ? WHERE id = ?', (curr_order, other_id))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@materials_bp.route('/api/material/asset/delete', methods=['POST'])
def delete_material_asset():
    data = request.json or {}
    mat_no = data.get('material_no', '').strip()
    asset_type = data.get('asset_type', '').strip()
    if not mat_no or not asset_type:
        return jsonify({'error': '参数不全'}), 400

    conn = get_db()
    cursor = conn.cursor()
    if asset_type == 'image':
        cursor.execute("UPDATE materials SET image_path = '' WHERE material_no = ?", (mat_no,))
    elif asset_type == 'manual':
        cursor.execute("UPDATE materials SET manual_path = '' WHERE material_no = ?", (mat_no,))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@materials_bp.route('/api/materials/export', methods=['GET'])
def export_materials_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    SELECT m.*,
           (SELECT COUNT(*) FROM material_structures s WHERE s.parent_material_no = m.material_no) as child_count,
           (SELECT GROUP_CONCAT(DISTINCT COALESCE(NULLIF(p.system_name, ''), NULLIF(p.code, ''), p.name))
            FROM bom_items b
            JOIN subsystems sub ON b.subsystem_id = sub.id
            JOIN projects p ON sub.project_id = p.id
            WHERE b.material_no = m.material_no) as referenced_projects
    FROM materials m
    ORDER BY m.material_no
    ''')
    rows = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "全局通用物料主库"

    headers = [
        "物料号", "元器件名称", "引用项目", "物料完整描述", "规格型号", "制造商",
        "单位", "图纸/代号", "子层级数", "实物图片", "技术手册(PDF)", "替代物料", "补充说明"
    ]
    ws.append(headers)

    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    from services.exporter import format_pipe_multiline
    data_font = Font(name="微软雅黑", size=10)
    for r_idx, r in enumerate(rows, start=2):
        row_values = [
            format_pipe_multiline(r['material_no']),
            format_pipe_multiline(r['name'] or ''),
            format_pipe_multiline(r['referenced_projects'] or '未引用'),
            format_pipe_multiline(r['description'] or ''),
            format_pipe_multiline(r['spec_model'] or ''),
            format_pipe_multiline(r['manufacturer'] or ''),
            format_pipe_multiline(r['unit'] or 'PC'),
            format_pipe_multiline(r['drawing_no'] or '/'),
            r['child_count'] if r['child_count'] > 0 else 0,
            "已配备" if r['image_path'] else "无",
            "已配备" if r['manual_path'] else "无",
            format_pipe_multiline(r['substitute'] or ''),
            format_pipe_multiline(r['notes'] or '')
        ]
        ws.append(row_values)
        max_lines = 1
        for c_idx in range(1, len(row_values) + 1):
            c = ws.cell(row=r_idx, column=c_idx)
            c.font = data_font
            c.border = thin_border
            if isinstance(c.value, str) and '\n' in c.value:
                lines_count = len(c.value.split('\n'))
                if lines_count > max_lines:
                    max_lines = lines_count
            if c_idx in [1, 7, 8, 9, 10, 11]:
                c.alignment = center_align
            else:
                c.alignment = left_align
        ws.row_dimensions[r_idx].height = max(22, max_lines * 18)

    col_widths = {
        1: 18, 2: 22, 3: 24, 4: 38, 5: 25, 6: 18,
        7: 8, 8: 16, 9: 12, 10: 12, 11: 14,
        12: 16, 13: 24
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="全局通用物料主库清单.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ----------------- Master Materials Excel Import API -----------------
@materials_bp.route('/api/materials/import/analyze', methods=['POST'])
def analyze_master_import():
    if 'file' not in request.files:
        return jsonify({'error': '未选择要上传的 Excel 文件'}), 400
    file = request.files['file']
    if not file.filename.lower().endswith('.xlsx'):
        return jsonify({'error': '请上传 .xlsx 格式的 Excel 文件'}), 400

    filename = secure_filename(file.filename) or 'master_import.xlsx'
    temp_filename = f"temp_master_{int(time.time())}_{filename}"
    save_path = os.path.join(UPLOAD_DIR, temp_filename)
    file.save(save_path)

    try:
        res = analyze_master_excel(save_path)
        res['temp_filename'] = temp_filename
        res['original_filename'] = file.filename
        return jsonify(res)
    except Exception as e:
        return jsonify({'error': f'解析 Excel 文件失败: {str(e)}'}), 500


@materials_bp.route('/api/materials/import/reanalyze', methods=['POST'])
def reanalyze_master_import():
    data = request.json or {}
    temp_filename = data.get('temp_filename', '').strip()
    sheet_name = data.get('sheet_name', '').strip()
    if not temp_filename:
        return jsonify({'error': '缺少临时文件名'}), 400

    save_path = os.path.join(UPLOAD_DIR, temp_filename)
    if not os.path.exists(save_path):
        return jsonify({'error': '上传的临时文件已过期，请重新选择文件上传'}), 400

    try:
        res = analyze_master_excel(save_path, sheet_name)
        res['temp_filename'] = temp_filename
        return jsonify(res)
    except Exception as e:
        return jsonify({'error': f'切换工作表分析失败: {str(e)}'}), 500


@materials_bp.route('/api/materials/import/execute', methods=['POST'])
def execute_master_import_api():
    data = request.json or {}
    temp_filename = data.get('temp_filename', '').strip()
    sheet_name = data.get('sheet_name', '').strip()
    header_row = data.get('header_row', 1)
    mapping = data.get('mapping', {})
    overwrite = data.get('overwrite', True)

    if not temp_filename:
        return jsonify({'error': '缺少临时文件名'}), 400

    save_path = os.path.join(UPLOAD_DIR, temp_filename)
    if not os.path.exists(save_path):
        return jsonify({'error': '上传的临时文件已过期，请重新选择文件上传'}), 400

    try:
        create_backup_snapshot(reason='import_master_excel', note=f'批量导入主库物料Excel [{temp_filename}] 前自动备份')
    except Exception as e_snap:
        logger.warning(f"Auto-backup before master import warning: {e_snap}")

    try:
        res = execute_master_import(save_path, sheet_name=sheet_name, header_row=header_row, mapping=mapping, overwrite=overwrite)
        if isinstance(res, dict) and 'error' in res:
            return jsonify(res), 400

        # 成功导入后自动清理 uploads/ 下的临时文件，防止磁盘堆积
        try:
            if os.path.exists(save_path):
                os.remove(save_path)
        except Exception:
            pass

        return jsonify(res)
    except Exception as e:
        return jsonify({'error': f'导入物料失败: {str(e)}'}), 500


@materials_bp.route('/api/material/<path:material_no>', methods=['GET'])
def get_single_material(material_no):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
    SELECT m.*,
           (SELECT COUNT(*) FROM material_structures s WHERE s.parent_material_no = m.material_no) as child_count,
           (SELECT COUNT(*) FROM bom_items b WHERE b.material_no = m.material_no) as usage_count
    FROM materials m
    WHERE m.material_no = ?
    ''', (material_no.strip(),))
    row = cursor.fetchone()
    conn.close()
    if row:
        return jsonify(dict(row))
    return jsonify(None), 404


# ----------------- Material Master Data Audit & De-duplication APIs -----------------
@materials_bp.route('/api/materials/audit', methods=['GET'])
def api_materials_audit():
    try:
        report = get_duplicate_materials_report()
        return jsonify({'success': True, 'report': report})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@materials_bp.route('/api/materials/check_duplicate', methods=['GET'])
def api_check_duplicate():
    name = request.args.get('name', '')
    spec_model = request.args.get('spec_model', '')
    mat_no = request.args.get('mat_no', '')
    try:
        matches = check_single_material_duplicate(name, spec_model, exclude_mat_no=mat_no)
        return jsonify({'success': True, 'matches': matches})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


# ----------------- AI Material Extraction API -----------------
@materials_bp.route('/api/materials/ai-parse', methods=['POST'])
def api_ai_parse_material():
    data = request.json or {}
    desc = (data.get('description') or '').strip()
    if not desc:
        return jsonify({'success': False, 'error': '物料描述不能为空'}), 400

    from services.ai_service import extract_material_semantics
    try:
        result = extract_material_semantics(desc)
        return jsonify(result)
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

