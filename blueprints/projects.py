# blueprints/projects.py
"""
Project lifecycle management, Subsystem Sheet CRUD, ECN revisions, Signatures,
BOM synchronization to Master, Project bundle save/open, Excel & Offline ZIP Exports.
"""

import os
import time
import sqlite3
from flask import Blueprint, request, jsonify, send_file
from werkzeug.utils import secure_filename
from services.database import get_db, find_template_path
from services.exporter import export_project_sbom
from services.offline_packager import generate_offline_package
from services.backup_manager import create_backup_snapshot
from services.project_bundle import (
    create_project_bundle_file,
    inspect_saved_project_file,
    import_project_from_file
)
from blueprints.common import (
    UPLOAD_DIR,
    EXPORT_DIR,
    check_project_unlocked
)

projects_bp = Blueprint('projects', __name__)


# ----------------- Project Query & Details -----------------
@projects_bp.route('/api/projects', methods=['GET'])
def get_projects():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM projects ORDER BY id DESC')
    projects = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return jsonify(projects)


@projects_bp.route('/api/project/<int:project_id>', methods=['GET'])
def get_project_detail(project_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM projects WHERE id = ?', (project_id,))
    project = cursor.fetchone()
    if not project:
        conn.close()
        return jsonify({'error': 'Project not found'}), 404

    project_dict = dict(project)

    # Subsystems
    cursor.execute('SELECT * FROM subsystems WHERE project_id = ? ORDER BY sort_order, id', (project_id,))
    project_dict['subsystems'] = [dict(s) for s in cursor.fetchall()]

    # Signatures
    cursor.execute('SELECT * FROM signatures WHERE project_id = ?', (project_id,))
    project_dict['signatures'] = [dict(s) for s in cursor.fetchall()]

    # ECN
    cursor.execute('SELECT * FROM ecn_records WHERE project_id = ? ORDER BY id', (project_id,))
    project_dict['ecn_records'] = [dict(e) for e in cursor.fetchall()]

    conn.close()
    return jsonify(project_dict)


# ----------------- Operator Identity Helper -----------------
def get_client_identity():
    """局域网内直接使用客户端 IP 地址作为协同身份标识。"""
    return request.headers.get('X-Forwarded-For', request.remote_addr or '127.0.0.1').split(',')[0].strip()



def check_and_expire_lease(cursor, project_row):
    """
    检查工程锁是否已过期（超过 180 秒无心跳）。
    若过期则自动置为锁定保护（释放占用），避免工程师关机/休眠后导致死锁。
    返回最新的 (is_locked, locked_by)。
    """
    import datetime
    if not project_row or project_row['is_locked'] == 1:
        return 1, ''
    locked_at_str = project_row['locked_at']
    if not locked_at_str:
        return 0, project_row['locked_by'] or ''
    try:
        t_str = str(locked_at_str).split('.')[0]
        locked_time = datetime.datetime.strptime(t_str, '%Y-%m-%d %H:%M:%S')
        now = datetime.datetime.now()
        if (now - locked_time).total_seconds() > 180:
            cursor.execute("UPDATE projects SET is_locked = 1, locked_by = '', locked_at = NULL WHERE id = ?", (project_row['id'],))
            return 1, ''
    except Exception:
        pass
    return project_row['is_locked'], project_row['locked_by'] or ''


# ----------------- Project Lock, Heartbeat & Unlock -----------------
@projects_bp.route('/api/project/<int:project_id>/heartbeat', methods=['POST'])
def heartbeat_project_api(project_id):
    """
    心跳续租接口：由持有锁的前端页面每 30 秒自动调用一次，
    刷新 locked_at 时间戳以维持工程编辑租约。
    """
    client_tag = get_client_identity()
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT id, is_locked, locked_by FROM projects WHERE id = ?', (project_id,))
        p = cursor.fetchone()
        if not p:
            return jsonify({'success': False, 'error': '工程不存在'}), 404

        if p['is_locked'] == 0:
            # 刷新租约
            cursor.execute("UPDATE projects SET locked_at = datetime('now', 'localtime') WHERE id = ?", (project_id,))
            conn.commit()
            return jsonify({'success': True, 'renewed': True, 'locked_by': p['locked_by']})
        return jsonify({'success': True, 'renewed': False, 'is_locked': p['is_locked']})
    finally:
        conn.close()


@projects_bp.route('/api/project/<int:project_id>/unlock', methods=['POST'])
def unlock_project_api(project_id):
    data = request.get_json(silent=True) or {}
    force = bool(data.get('force', False))
    client_tag = get_client_identity()

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT id, code, system_name, name, is_locked, locked_by, locked_at FROM projects WHERE id = ?', (project_id,))
        p = cursor.fetchone()
        if not p:
            return jsonify({'success': False, 'error': '工程不存在或已被删除'}), 404

        p_name = p['system_name'] or p['code'] or p['name']
        curr_locked, curr_user = check_and_expire_lease(cursor, p)
        curr_time = p['locked_at'] or ''

        # 若工程当前已解锁，且被其他同事占用，且不是强制接管请求
        if curr_locked == 0 and curr_user and curr_user != client_tag and not force:
            return jsonify({
                'success': False,
                'need_confirm': True,
                'locked_by': curr_user,
                'locked_at': curr_time,
                'error': f'工程【{p_name}】当前正由同事 [{curr_user}] 于 {curr_time} 解锁占用编辑中！\n是否确认强制接管该工程的编辑权限？'
            }), 409

        # 解锁当前工程并登记占用人（独立项目锁：完全不影响其他工程状态！）
        cursor.execute('''
            UPDATE projects 
            SET is_locked = 0, locked_by = ?, locked_at = datetime('now', 'localtime')
            WHERE id = ?
        ''', (client_tag, project_id))
        conn.commit()

        return jsonify({
            'success': True,
            'project_id': project_id,
            'is_locked': 0,
            'locked_by': client_tag,
            'message': f'工程【{p_name}】已成功解锁并激活编辑状态！'
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'解锁工程失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass



@projects_bp.route('/api/project/<int:project_id>/lock', methods=['POST'])
def lock_project_api(project_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT id, code, system_name, name FROM projects WHERE id = ?', (project_id,))
        p = cursor.fetchone()
        if not p:
            return jsonify({'success': False, 'error': '工程不存在或已被删除'}), 404

        # 锁定当前工程并清空占用人（完全不影响其他工程）
        cursor.execute("UPDATE projects SET is_locked = 1, locked_by = '', locked_at = NULL WHERE id = ?", (project_id,))
        conn.commit()

        p_name = p['system_name'] or p['code'] or p['name']
        return jsonify({
            'success': True,
            'project_id': project_id,
            'is_locked': 1,
            'locked_by': '',
            'message': f'工程【{p_name}】已设置为锁定保护状态（只读）！'
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'锁定工程失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


# ----------------- Project Save, Rename, Copy & Delete -----------------
@projects_bp.route('/api/project/save', methods=['POST'])
def save_project():
    data = request.get_json(silent=True) or {}
    project_id = data.get('id')

    code = (data.get('code') or '').strip()
    name = (data.get('name') or '').strip()
    system_name = (data.get('system_name') or '').strip()

    if not code and not system_name and not name:
        return jsonify({'success': False, 'error': '项目名称不能为空'}), 400

    if not name:
        name = '风力发电机组'
    if not system_name:
        system_name = code if code else (name if name != '风力发电机组' else '未命名项目')
    if not code:
        code = system_name if system_name else 'EW-DEFAULT'

    if project_id:
        is_ok, err_msg = check_project_unlocked(project_id=project_id)
        if not is_ok:
            return jsonify({'success': False, 'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    try:
        if project_id:
            cursor.execute('SELECT id FROM projects WHERE code = ? AND id != ?', (code, project_id))
            if cursor.fetchone():
                return jsonify({'success': False, 'error': f'机型/项目代号 "{code}" 已被其他项目占用，请更换代号'}), 400

            cursor.execute('''
            UPDATE projects SET
                code=?, name=?, system_name=?, doc_no=?, ecn_no=?, form_no=?, company=?
            WHERE id=?
            ''', (code, name, system_name,
                  data.get('doc_no', ''), data.get('ecn_no', ''),
                  data.get('form_no', 'FDJL-JS-037'), data.get('company', '上海电气风电集团股份有限公司'),
                  project_id))
        else:
            base_code = code
            suffix = 1
            inserted = False
            while not inserted:
                candidate_code = base_code if suffix == 1 else f"{base_code}_{suffix}"
                cursor.execute('SELECT id FROM projects WHERE code = ?', (candidate_code,))
                if cursor.fetchone():
                    suffix += 1
                    continue
                try:
                    client_tag = get_client_identity()
                    cursor.execute('''
                    INSERT INTO projects (code, name, system_name, doc_no, ecn_no, form_no, company, is_locked, locked_by, locked_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?, datetime('now', 'localtime'))
                    ''', (candidate_code, name, system_name,
                          data.get('doc_no', ''), data.get('ecn_no', ''),
                          data.get('form_no', 'FDJL-JS-037'), data.get('company', '上海电气风电集团股份有限公司'),
                          client_tag))
                    project_id = cursor.lastrowid
                    code = candidate_code
                    inserted = True
                except sqlite3.IntegrityError:
                    suffix += 1

            default_subs = ['SBOM']
            tpl_path = find_template_path()
            if tpl_path and os.path.exists(tpl_path):
                try:
                    import openpyxl
                    wb_tpl = openpyxl.load_workbook(tpl_path, read_only=True)
                    candidates = [s for s in wb_tpl.sheetnames if '封面' not in s and '签字' not in s and '修改' not in s]
                    if candidates:
                        default_subs = candidates
                    wb_tpl.close()
                except Exception:
                    pass
            for idx, sub_name in enumerate(default_subs):
                cursor.execute('INSERT INTO subsystems (project_id, name, sort_order) VALUES (?, ?, ?)',
                               (project_id, sub_name, idx))

        conn.commit()
        return jsonify({'success': True, 'id': project_id, 'name': name, 'system_name': system_name, 'code': code})
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'保存项目失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@projects_bp.route('/api/project/rename', methods=['POST'])
def rename_project():
    data = request.json or {}
    project_id = data.get('id')
    new_system_name = (data.get('system_name') or data.get('name') or '').strip()
    new_code = (data.get('code') or '').strip()

    if not project_id or not new_system_name:
        return jsonify({'error': '缺少项目ID或新名称'}), 400

    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT * FROM projects WHERE id = ?', (project_id,))
        proj = cursor.fetchone()
        if not proj:
            return jsonify({'success': False, 'error': '项目不存在'}), 404

        custom_product_name = (data.get('product_name') or '').strip()
        final_name = custom_product_name or proj['name'] or '风力发电机组'

        final_code = new_code or proj['code']
        cursor.execute('SELECT id FROM projects WHERE code = ? AND id != ?', (final_code, project_id))
        if cursor.fetchone():
            return jsonify({'success': False, 'error': f'代号 "{final_code}" 已被其他项目占用，请指定其他代号'}), 400

        cursor.execute('''
        UPDATE projects SET name = ?, system_name = ?, code = ? WHERE id = ?
        ''', (final_name, new_system_name, final_code, project_id))
        conn.commit()
        return jsonify({'success': True, 'id': project_id, 'name': final_name, 'system_name': new_system_name, 'code': final_code})
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'重命名项目失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@projects_bp.route('/api/project/copy', methods=['POST'])
def copy_project():
    data = request.json or {}
    project_id = data.get('id')
    if not project_id:
        return jsonify({'error': '缺少原项目ID'}), 400

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT * FROM projects WHERE id = ?', (project_id,))
        orig = cursor.fetchone()
        if not orig:
            return jsonify({'success': False, 'error': '原项目不存在或已被删除'}), 404

        user_name = (data.get('name') or '').strip()
        user_system_name = (data.get('system_name') or '').strip()
        user_code = (data.get('code') or '').strip()

        orig_name = orig['name'] or '风力发电机组'
        orig_system_name = orig['system_name'] or ''
        orig_code = orig['code'] or ''

        user_name = orig_name or '风力发电机组'

        if not user_system_name:
            user_system_name = f"{orig_system_name}_副本" if orig_system_name else f"{orig_code}_副本"

        if not user_code:
            base_code = f"{orig_code}_copy"
            candidate = base_code
            suffix = 2
            while True:
                cursor.execute('SELECT id FROM projects WHERE code = ?', (candidate,))
                if not cursor.fetchone():
                    user_code = candidate
                    break
                candidate = f"{base_code}_{suffix}"
                suffix += 1
        else:
            cursor.execute('SELECT id FROM projects WHERE code = ?', (user_code,))
            if cursor.fetchone():
                return jsonify({'success': False, 'error': f'机型代号 "{user_code}" 已被占用，请指定新的代号'}), 400

        cursor.execute('''
            INSERT INTO projects (code, name, system_name, doc_no, ecn_no, form_no, company)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        ''', (user_code, user_name, user_system_name,
              orig['doc_no'] or '', orig['ecn_no'] or '',
              orig['form_no'] or 'FDJL-JS-037', orig['company'] or '上海电气风电集团股份有限公司'))
        new_project_id = cursor.lastrowid

        cursor.execute('SELECT * FROM subsystems WHERE project_id = ? ORDER BY sort_order, id', (project_id,))
        subsystems = cursor.fetchall()
        for sub in subsystems:
            cursor.execute('''
                INSERT INTO subsystems (project_id, name, system_title, doc_no, sort_order)
                VALUES (?, ?, ?, ?, ?)
            ''', (new_project_id, sub['name'], sub['system_title'] or '', sub['doc_no'] or '', sub['sort_order']))
            new_sub_id = cursor.lastrowid

            cursor.execute('''
                SELECT row_num, level_str, level_num, seq_num, material_no, drawing_no, qty, unit, notes, substitute
                FROM bom_items
                WHERE subsystem_id = ?
                ORDER BY row_num, id
            ''', (sub['id'],))
            items = cursor.fetchall()
            for it in items:
                cursor.execute('''
                    INSERT INTO bom_items (subsystem_id, row_num, level_str, level_num, seq_num, material_no, drawing_no, qty, unit, notes, substitute)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (new_sub_id, it['row_num'], it['level_str'], it['level_num'], it['seq_num'],
                      it['material_no'], it['drawing_no'], it['qty'], it['unit'], it['notes'], it['substitute']))

        cursor.execute('SELECT role, sign_name, sign_date FROM signatures WHERE project_id = ?', (project_id,))
        signatures = cursor.fetchall()
        for s in signatures:
            cursor.execute('''
                INSERT INTO signatures (project_id, role, sign_name, sign_date)
                VALUES (?, ?, ?, ?)
            ''', (new_project_id, s['role'], s['sign_name'] or '', s['sign_date'] or ''))

        cursor.execute('SELECT ecn_no, change_desc, changer, change_date FROM ecn_records WHERE project_id = ? ORDER BY id', (project_id,))
        ecn_records = cursor.fetchall()
        for e in ecn_records:
            cursor.execute('''
                INSERT INTO ecn_records (project_id, ecn_no, change_desc, changer, change_date)
                VALUES (?, ?, ?, ?, ?)
            ''', (new_project_id, e['ecn_no'] or '', e['change_desc'] or '', e['changer'] or '', e['change_date'] or ''))

        client_tag = get_client_identity()
        cursor.execute("UPDATE projects SET is_locked = 0, locked_by = ?, locked_at = datetime('now', 'localtime') WHERE id = ?", (client_tag, new_project_id))

        conn.commit()
        return jsonify({
            'success': True,
            'id': new_project_id,
            'name': user_name,
            'system_name': user_system_name,
            'code': user_code,
            'message': f'项目已成功完整复制为【{user_system_name}】！'
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'复制项目失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@projects_bp.route('/api/project/delete', methods=['POST'])
def delete_project():
    data = request.json or {}
    project_id = data.get('id')
    if not project_id:
        return jsonify({'error': '缺少项目ID'}), 400

    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT id, name, code, system_name FROM projects WHERE id = ?', (project_id,))
        proj = cursor.fetchone()
        if not proj:
            return jsonify({'success': False, 'error': '项目不存在或已被删除'}), 404

        proj_name = proj['name']
        proj_code = proj['code']

        try:
            create_backup_snapshot(reason='delete_project', note=f'删除项目 [{proj_code} - {proj_name}] 前自动备份')
        except Exception as e_snap:
            print("Auto-backup before project delete warning:", e_snap)

        cursor.execute('DELETE FROM bom_items WHERE subsystem_id IN (SELECT id FROM subsystems WHERE project_id = ?)', (project_id,))
        cursor.execute('DELETE FROM subsystems WHERE project_id = ?', (project_id,))
        cursor.execute('DELETE FROM signatures WHERE project_id = ?', (project_id,))
        cursor.execute('DELETE FROM ecn_records WHERE project_id = ?', (project_id,))
        cursor.execute('DELETE FROM projects WHERE id = ?', (project_id,))

        conn.commit()
        return jsonify({
            'success': True,
            'deleted_id': project_id,
            'deleted_name': proj_name,
            'deleted_code': proj_code,
            'message': f"项目【{proj_code} {proj_name}】已成功永久删除！"
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'删除项目失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


# ----------------- Project ECN & Signatures APIs -----------------
@projects_bp.route('/api/project/<int:project_id>/ecn', methods=['GET', 'POST'])
def project_ecn_api(project_id):
    if request.method == 'GET':
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM ecn_records WHERE project_id = ? ORDER BY id', (project_id,))
        records = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return jsonify({'success': True, 'records': records})

    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    data = request.get_json(silent=True) or {}
    ecn_no = (data.get('ecn_no') or '').strip()
    change_desc = (data.get('change_desc') or '').strip()
    changer = (data.get('changer') or '').strip()
    change_date = (data.get('change_date') or '').strip()

    if not ecn_no and not change_desc:
        return jsonify({'success': False, 'error': 'ECN号和更改内容描述不能同时为空'}), 400

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
        INSERT INTO ecn_records (project_id, ecn_no, change_desc, changer, change_date)
        VALUES (?, ?, ?, ?, ?)
        ''', (project_id, ecn_no, change_desc, changer, change_date))
        new_id = cursor.lastrowid
        conn.commit()
        return jsonify({
            'success': True,
            'message': '新增修改记录成功',
            'record': {
                'id': new_id,
                'project_id': project_id,
                'ecn_no': ecn_no,
                'change_desc': change_desc,
                'changer': changer,
                'change_date': change_date
            }
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'添加失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@projects_bp.route('/api/project/<int:project_id>/ecn/<int:ecn_id>', methods=['PUT', 'DELETE'])
def project_ecn_item_api(project_id, ecn_id):
    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT id FROM ecn_records WHERE id = ? AND project_id = ?', (ecn_id, project_id))
    if not cursor.fetchone():
        conn.close()
        return jsonify({'success': False, 'error': '未找到该ECN记录'}), 404

    try:
        if request.method == 'DELETE':
            cursor.execute('DELETE FROM ecn_records WHERE id = ?', (ecn_id,))
            conn.commit()
            return jsonify({'success': True, 'message': '修改记录已删除'})

        data = request.get_json(silent=True) or {}
        ecn_no = (data.get('ecn_no') or '').strip()
        change_desc = (data.get('change_desc') or '').strip()
        changer = (data.get('changer') or '').strip()
        change_date = (data.get('change_date') or '').strip()

        cursor.execute('''
        UPDATE ecn_records SET
            ecn_no = ?, change_desc = ?, changer = ?, change_date = ?
        WHERE id = ?
        ''', (ecn_no, change_desc, changer, change_date, ecn_id))
        conn.commit()
        return jsonify({'success': True, 'message': '修改记录已更新'})
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'操作失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


@projects_bp.route('/api/project/<int:project_id>/signatures/save', methods=['POST'])
def save_project_signatures(project_id):
    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    data = request.get_json(silent=True) or {}
    signatures = data.get('signatures') or []

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('DELETE FROM signatures WHERE project_id = ?', (project_id,))
        for s in signatures:
            role = (s.get('role') or '').strip()
            sign_name = (s.get('sign_name') or '').strip()
            sign_date = (s.get('sign_date') or '').strip()
            if role:
                cursor.execute('''
                INSERT INTO signatures (project_id, role, sign_name, sign_date)
                VALUES (?, ?, ?, ?)
                ''', (project_id, role, sign_name, sign_date))
        conn.commit()
        return jsonify({'success': True, 'message': '签字信息保存成功'})
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'success': False, 'error': f'保存签字失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass





# ----------------- Project Pull / Refresh from Master Library -----------------
@projects_bp.route('/api/project/<int:project_id>/check_master_sync', methods=['GET'])
def check_master_sync_api(project_id):
    """
    Scans all materials referenced in current project and checks for updates/differences
    available in the Master Materials Library (e.g. photos, manuals, drawings, substitutes, units).
    """
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT id, name, code, is_locked FROM projects WHERE id = ?', (project_id,))
    proj = cursor.fetchone()
    if not proj:
        conn.close()
        return jsonify({'success': False, 'error': '项目不存在'}), 404

    cursor.execute('''
    SELECT b.id, b.subsystem_id, s.name as subsystem_name, b.row_num, b.material_no,
           b.drawing_no as b_drawing_no, b.unit as b_unit, b.notes as b_notes, b.substitute as b_substitute
    FROM bom_items b
    JOIN subsystems s ON b.subsystem_id = s.id
    WHERE s.project_id = ?
    ORDER BY s.sort_order, s.id, b.row_num, b.id
    ''', (project_id,))
    bom_rows = [dict(r) for r in cursor.fetchall()]

    mat_nos = list(dict.fromkeys([(r['material_no'] or '').strip() for r in bom_rows if (r['material_no'] or '').strip()]))

    proj_mat_map = {}
    for r in bom_rows:
        mno = (r['material_no'] or '').strip()
        if not mno:
            continue
        if mno not in proj_mat_map:
            proj_mat_map[mno] = {
                'occurrences': 0,
                'subsystems': set(),
                'rows': []
            }
        proj_mat_map[mno]['occurrences'] += 1
        if r.get('subsystem_name'):
            proj_mat_map[mno]['subsystems'].add(r['subsystem_name'])
        proj_mat_map[mno]['rows'].append(r)

    master_map = {}
    if mat_nos:
        placeholders = ','.join(['?'] * len(mat_nos))
        cursor.execute(f'''
        SELECT material_no, name, description, spec_model, manufacturer, unit, drawing_no, notes, substitute, image_path, manual_path
        FROM materials
        WHERE material_no IN ({placeholders})
        ''', mat_nos)
        master_map = {r['material_no']: dict(r) for r in cursor.fetchall()}

    conn.close()

    items = []
    has_update_count = 0
    has_photo_count = 0
    has_manual_count = 0
    missing_in_master_count = 0

    for mno in mat_nos:
        p_info = proj_mat_map[mno]
        m_mat = master_map.get(mno)

        if not m_mat:
            missing_in_master_count += 1
            items.append({
                'material_no': mno,
                'name': '',
                'description': '',
                'spec_model': '',
                'manufacturer': '',
                'status': 'not_in_master',
                'status_label': '未入主库',
                'has_updates': False,
                'diff_types': [],
                'diff_summary': '该物料在全局主库中尚无记录（可先同步主库）',
                'subsystems': list(p_info['subsystems']),
                'occurrences': p_info['occurrences'],
                'has_photo': False,
                'has_manual': False,
                'image_path': '',
                'manual_path': '',
                'm_drawing_no': '',
                'curr_drawing_no': p_info['rows'][0]['b_drawing_no'] if p_info['rows'] else '/',
                'm_substitute': '',
                'curr_substitute': p_info['rows'][0]['b_substitute'] if p_info['rows'] else '',
                'm_unit': '',
                'curr_unit': p_info['rows'][0]['b_unit'] if p_info['rows'] else 'PCS',
                'm_notes': '',
                'curr_notes': p_info['rows'][0]['b_notes'] if p_info['rows'] else ''
            })
            continue

        diff_types = []
        diff_reasons = []

        m_dwg = (m_mat.get('drawing_no') or '').strip()
        m_unit = (m_mat.get('unit') or '').strip()
        m_sub = (m_mat.get('substitute') or '').strip()
        m_notes = (m_mat.get('notes') or '').strip()
        has_photo = bool((m_mat.get('image_path') or '').strip())
        has_manual = bool((m_mat.get('manual_path') or '').strip())

        if has_photo:
            has_photo_count += 1
        if has_manual:
            has_manual_count += 1

        # Check drawing_no diff across project rows
        drawing_diff_found = False
        if m_dwg and m_dwg != '/':
            for r in p_info['rows']:
                b_dwg = (r['b_drawing_no'] or '').strip()
                if not b_dwg or b_dwg == '/' or b_dwg != m_dwg:
                    drawing_diff_found = True
                    break
        if drawing_diff_found:
            diff_types.append('drawing_no')
            diff_reasons.append(f'主库标准图号: {m_dwg}')

        # Check substitute diff
        sub_diff_found = False
        if m_sub:
            for r in p_info['rows']:
                b_sub = (r['b_substitute'] or '').strip()
                if not b_sub or b_sub != m_sub:
                    sub_diff_found = True
                    break
        if sub_diff_found:
            diff_types.append('substitute')
            diff_reasons.append(f'主库推荐替代料: {m_sub}')

        # Check unit diff
        unit_diff_found = False
        if m_unit:
            for r in p_info['rows']:
                b_u = (r['b_unit'] or '').strip()
                if b_u and b_u != m_unit:
                    unit_diff_found = True
                    break
        if unit_diff_found:
            diff_types.append('unit')
            diff_reasons.append(f'主库标准单位: {m_unit}')

        # Check notes diff
        notes_diff_found = False
        if m_notes:
            for r in p_info['rows']:
                b_n = (r['b_notes'] or '').strip()
                if not b_n:
                    notes_diff_found = True
                    break
        if notes_diff_found:
            diff_types.append('notes')
            diff_reasons.append('主库含技术说明')

        has_updates = len(diff_types) > 0
        if has_updates:
            has_update_count += 1

        curr_dwg = p_info['rows'][0]['b_drawing_no'] if p_info['rows'] else '/'
        curr_sub = p_info['rows'][0]['b_substitute'] if p_info['rows'] else ''
        curr_unit = p_info['rows'][0]['b_unit'] if p_info['rows'] else 'PCS'
        curr_notes = p_info['rows'][0]['b_notes'] if p_info['rows'] else ''

        items.append({
            'material_no': mno,
            'name': m_mat.get('name') or '',
            'description': m_mat.get('description') or '',
            'spec_model': m_mat.get('spec_model') or '',
            'manufacturer': m_mat.get('manufacturer') or '',
            'status': 'has_diff' if has_updates else 'aligned',
            'status_label': '可更新' if has_updates else '已对齐',
            'has_updates': has_updates,
            'diff_types': diff_types,
            'diff_summary': '；'.join(diff_reasons) if diff_reasons else '当前项目数据与主库完全一致',
            'subsystems': list(p_info['subsystems']),
            'occurrences': p_info['occurrences'],
            'has_photo': has_photo,
            'has_manual': has_manual,
            'image_path': m_mat.get('image_path') or '',
            'manual_path': m_mat.get('manual_path') or '',
            'm_drawing_no': m_dwg,
            'curr_drawing_no': curr_dwg,
            'm_substitute': m_sub,
            'curr_substitute': curr_sub,
            'm_unit': m_unit,
            'curr_unit': curr_unit,
            'm_notes': m_notes,
            'curr_notes': curr_notes
        })

    # Sort items: has_updates first, then missing_in_master, then aligned
    def sort_key(it):
        if it['has_updates']:
            return (0, it['material_no'])
        if it['status'] == 'not_in_master':
            return (1, it['material_no'])
        return (2, it['material_no'])

    items.sort(key=sort_key)

    return jsonify({
        'success': True,
        'project_id': project_id,
        'project_code': proj['code'] or '',
        'project_name': proj['name'] or '',
        'is_locked': proj['is_locked'],
        'total_materials': len(mat_nos),
        'has_update_count': has_update_count,
        'has_photo_count': has_photo_count,
        'has_manual_count': has_manual_count,
        'missing_in_master_count': missing_in_master_count,
        'items': items
    })


@projects_bp.route('/api/project/<int:project_id>/pull_from_master', methods=['POST'])
def pull_from_master_api(project_id):
    """
    Applies selected updates from the Master Materials Library into the current project's BOM items.
    """
    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'success': False, 'error': err_msg}), 403

    data = request.json or {}
    selected_mat_nos = data.get('material_nos')
    sync_options = data.get('sync_options') or {
        'sync_drawing_no': True,
        'sync_unit': True,
        'sync_substitute': True,
        'sync_notes': True
    }

    if not selected_mat_nos:
        return jsonify({'success': False, 'error': '未选择任何需要拉取更新的物料'}), 400

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute('SELECT id, name, code FROM projects WHERE id = ?', (project_id,))
    proj = cursor.fetchone()
    if not proj:
        conn.close()
        return jsonify({'success': False, 'error': '项目不存在'}), 404

    proj_title = (proj['code'] or '') + (' ' + proj['name'] if proj['name'] else '')

    # Automatic safety snapshot before pulling from master
    try:
        create_backup_snapshot(
            reason='sync_from_master',
            note=f'从主库拉取物料更新到项目 [{proj_title.strip()}] 前自动备份'
        )
    except Exception as e_snap:
        print("Auto-backup before pull_from_master warning:", e_snap)

    placeholders = ','.join(['?'] * len(selected_mat_nos))
    cursor.execute(f'''
    SELECT material_no, drawing_no, unit, notes, substitute
    FROM materials
    WHERE material_no IN ({placeholders})
    ''', selected_mat_nos)
    master_rows = {r['material_no']: dict(r) for r in cursor.fetchall()}

    cursor.execute('''
    SELECT b.id, b.material_no, b.drawing_no, b.unit, b.notes, b.substitute
    FROM bom_items b
    JOIN subsystems s ON b.subsystem_id = s.id
    WHERE s.project_id = ?
    ''', (project_id,))
    bom_items = cursor.fetchall()

    updated_rows_count = 0
    updated_mats = set()

    for it in bom_items:
        mno = (it['material_no'] or '').strip()
        if mno not in master_rows:
            continue

        m_mat = master_rows[mno]
        updates = []
        params = []

        if sync_options.get('sync_drawing_no', True):
            m_dwg = (m_mat.get('drawing_no') or '').strip()
            curr_dwg = (it['drawing_no'] or '').strip()
            if m_dwg and m_dwg != '/' and (not curr_dwg or curr_dwg == '/' or curr_dwg != m_dwg):
                updates.append("drawing_no = ?")
                params.append(m_dwg)

        if sync_options.get('sync_unit', True):
            m_u = (m_mat.get('unit') or '').strip()
            curr_u = (it['unit'] or '').strip()
            if m_u and (not curr_u or curr_u != m_u):
                updates.append("unit = ?")
                params.append(m_u)

        if sync_options.get('sync_substitute', True):
            m_sub = (m_mat.get('substitute') or '').strip()
            curr_sub = (it['substitute'] or '').strip()
            if m_sub and (not curr_sub or curr_sub != m_sub):
                updates.append("substitute = ?")
                params.append(m_sub)

        if sync_options.get('sync_notes', True):
            m_notes = (m_mat.get('notes') or '').strip()
            curr_notes = (it['notes'] or '').strip()
            if m_notes and not curr_notes:
                updates.append("notes = ?")
                params.append(m_notes)

        if updates:
            params.append(it['id'])
            cursor.execute(f"UPDATE bom_items SET {', '.join(updates)} WHERE id = ?", params)
            updated_rows_count += 1
            updated_mats.add(mno)

    conn.commit()
    conn.close()

    summary_msg = f"成功从物料主库拉取 {len(updated_mats)} 种物料的标准数据，共更新了 {updated_rows_count} 处BOM明细行！"
    return jsonify({
        'success': True,
        'updated_materials_count': len(updated_mats),
        'updated_rows_count': updated_rows_count,
        'message': summary_msg
    })


# ----------------- Project Save & Open (.sbom / .json) API -----------------
@projects_bp.route('/api/project/<int:project_id>/export_file', methods=['GET'])
def export_project_file_api(project_id):
    bundle_format = request.args.get('format', 'sbom').lower()
    if bundle_format not in ['sbom', 'json', 'zip']:
        bundle_format = 'sbom'

    buffer, filename, mimetype = create_project_bundle_file(project_id, bundle_format)
    if not buffer:
        return jsonify({'error': '项目未找到或导出失败'}), 404

    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype=mimetype
    )


@projects_bp.route('/api/project/open/inspect', methods=['POST'])
def inspect_open_project_api():
    if 'file' not in request.files:
        return jsonify({'error': '未提供文件'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'error': '文件名为空'}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ['.sbom', '.json', '.zip']:
        return jsonify({'error': '请选择 .sbom、.json 或 .zip 格式的已保存项目文件'}), 400

    filename = secure_filename(file.filename) or 'saved_project.sbom'
    temp_filename = f"temp_open_{int(time.time())}_{filename}"
    save_path = os.path.join(UPLOAD_DIR, temp_filename)
    file.save(save_path)

    try:
        res = inspect_saved_project_file(save_path)
        res['temp_filename'] = temp_filename
        res['original_filename'] = file.filename
        return jsonify(res)
    except Exception as e:
        return jsonify({'error': f'解析项目文件失败: {str(e)}'}), 400


@projects_bp.route('/api/project/open/execute', methods=['POST'])
def execute_open_project_api():
    data = request.json or {}
    temp_filename = data.get('temp_filename', '').strip()
    action = data.get('action', 'new').strip()
    target_project_id = data.get('target_project_id')

    if not temp_filename:
        return jsonify({'error': '缺少临时文件名'}), 400

    save_path = os.path.join(UPLOAD_DIR, temp_filename)
    if not os.path.exists(save_path):
        return jsonify({'error': '上传的工程文件已失效，请重新选择文件打开'}), 400

    try:
        create_backup_snapshot(reason='open_project_file', note=f'载入/恢复工程文件 [{temp_filename}] 前自动备份')
    except Exception as e_snap:
        print("Auto-backup before open project file warning:", e_snap)

    try:
        res = import_project_from_file(save_path, action=action, target_project_id=target_project_id)
        if res and res.get('success') and res.get('id'):
            opened_id = res.get('id')
            conn = get_db()
            cursor = conn.cursor()
            client_tag = get_client_identity()
            cursor.execute("UPDATE projects SET is_locked = 0, locked_by = ?, locked_at = datetime('now', 'localtime') WHERE id = ?", (client_tag, opened_id))
            conn.commit()
            conn.close()
        return jsonify(res)
    except Exception as e:
        return jsonify({'error': f'载入项目失败: {str(e)}'}), 500


# ----------------- Subsystems Sheet Management -----------------
@projects_bp.route('/api/subsystem/add', methods=['POST'])
def add_subsystem():
    data = request.json
    project_id = data.get('project_id')
    name = data.get('name', '').strip()
    if not project_id or not name:
        return jsonify({'error': '缺少项目ID或Sheet名称'}), 400

    is_ok, err_msg = check_project_unlocked(project_id=project_id)
    if not is_ok:
        return jsonify({'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT COALESCE(MAX(sort_order), 0) + 1 FROM subsystems WHERE project_id = ?', (project_id,))
    next_order = cursor.fetchone()[0]
    cursor.execute('''
    INSERT INTO subsystems (project_id, name, sort_order)
    VALUES (?, ?, ?)
    ''', (project_id, name, next_order))
    new_id = cursor.lastrowid
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'id': new_id, 'name': name})


@projects_bp.route('/api/subsystem/delete', methods=['POST'])
def delete_subsystem():
    data = request.json
    subsystem_id = data.get('id')
    if not subsystem_id:
        return jsonify({'error': '缺少Sheet ID'}), 400

    is_ok, err_msg = check_project_unlocked(subsystem_id=subsystem_id)
    if not is_ok:
        return jsonify({'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT project_id, name FROM subsystems WHERE id = ?', (subsystem_id,))
    sub = cursor.fetchone()
    if not sub:
        conn.close()
        return jsonify({'error': 'Sheet不存在'}), 404

    project_id = sub[0]
    cursor.execute('SELECT count(*) FROM subsystems WHERE project_id = ?', (project_id,))
    count = cursor.fetchone()[0]
    if count <= 1:
        conn.close()
        return jsonify({'error': '至少需要保留一个BOM Sheet，无法删除最后一个！'}), 400

    cursor.execute('DELETE FROM bom_items WHERE subsystem_id = ?', (subsystem_id,))
    cursor.execute('DELETE FROM subsystems WHERE id = ?', (subsystem_id,))
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'project_id': project_id})


@projects_bp.route('/api/subsystem/rename', methods=['POST'])
def rename_subsystem():
    data = request.json
    subsystem_id = data.get('id')
    new_name = data.get('name', '').strip()
    if not subsystem_id or not new_name:
        return jsonify({'error': '缺少参数'}), 400

    is_ok, err_msg = check_project_unlocked(subsystem_id=subsystem_id)
    if not is_ok:
        return jsonify({'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('UPDATE subsystems SET name = ? WHERE id = ?', (new_name, subsystem_id))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@projects_bp.route('/api/subsystem/copy', methods=['POST'])
def copy_subsystem():
    data = request.json or {}
    subsystem_id = data.get('id')
    new_name = (data.get('name') or '').strip()

    if not subsystem_id:
        return jsonify({'error': '缺少Sheet ID'}), 400

    is_ok, err_msg = check_project_unlocked(subsystem_id=subsystem_id)
    if not is_ok:
        return jsonify({'error': err_msg}), 403

    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('SELECT project_id, name, system_title, doc_no, sort_order FROM subsystems WHERE id = ?', (subsystem_id,))
        sub = cursor.fetchone()
        if not sub:
            return jsonify({'error': '源Sheet不存在'}), 404

        project_id = sub['project_id']
        orig_name = sub['name']

        if not new_name:
            base_name = f"{orig_name} - 副本"
            candidate_name = base_name
            suffix = 2
            cursor.execute('SELECT name FROM subsystems WHERE project_id = ?', (project_id,))
            existing_names = {r['name'] for r in cursor.fetchall()}
            while candidate_name in existing_names:
                candidate_name = f"{base_name} ({suffix})"
                suffix += 1
            new_name = candidate_name
        else:
            cursor.execute('SELECT id FROM subsystems WHERE project_id = ? AND name = ?', (project_id, new_name))
            if cursor.fetchone():
                return jsonify({'error': f'Sheet名称【{new_name}】已存在，请使用其他名称'}), 400

        current_sort_order = sub['sort_order'] if sub['sort_order'] is not None else 0
        new_sort_order = current_sort_order + 1
        cursor.execute('''
            UPDATE subsystems 
            SET sort_order = sort_order + 1 
            WHERE project_id = ? AND sort_order >= ?
        ''', (project_id, new_sort_order))

        cursor.execute('''
            INSERT INTO subsystems (project_id, name, system_title, doc_no, sort_order)
            VALUES (?, ?, ?, ?, ?)
        ''', (project_id, new_name, sub['system_title'] or '', sub['doc_no'] or '', new_sort_order))
        new_subsystem_id = cursor.lastrowid

        cursor.execute('''
            SELECT row_num, level_str, level_num, seq_num, material_no, drawing_no, qty, unit, notes, substitute
            FROM bom_items
            WHERE subsystem_id = ?
            ORDER BY row_num, id
        ''', (subsystem_id,))
        items = cursor.fetchall()

        for it in items:
            cursor.execute('''
                INSERT INTO bom_items (subsystem_id, row_num, level_str, level_num, seq_num, material_no, drawing_no, qty, unit, notes, substitute)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (new_subsystem_id, it['row_num'], it['level_str'], it['level_num'], it['seq_num'],
                  it['material_no'], it['drawing_no'], it['qty'], it['unit'], it['notes'], it['substitute']))

        conn.commit()
        return jsonify({
            'success': True,
            'new_id': new_subsystem_id,
            'name': new_name,
            'project_id': project_id,
            'item_count': len(items),
            'message': f'已成功复制 Sheet【{new_name}】，包含 {len(items)} 条物料记录！'
        })
    except Exception as e:
        try:
            conn.rollback()
        except Exception:
            pass
        return jsonify({'error': f'复制Sheet失败: {str(e)}'}), 500
    finally:
        try:
            conn.close()
        except Exception:
            pass


# ----------------- Project Exports (Excel SBOM & Offline Field ZIP) -----------------
@projects_bp.route('/api/export/<int:project_id>', methods=['GET'])
def export_excel(project_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT code, name, system_name FROM projects WHERE id = ?', (project_id,))
    proj = cursor.fetchone()
    conn.close()
    if not proj:
        return jsonify({'error': 'Project not found'}), 404

    system_name = (proj['system_name'] or '').strip()
    name = (proj['name'] or '').strip()
    code = (proj['code'] or '').strip()

    if system_name and system_name != '风力发电机组':
        raw_name = system_name
    elif name and name != '风力发电机组':
        raw_name = name
    elif system_name:
        raw_name = system_name
    elif code:
        raw_name = code
    else:
        raw_name = 'Service_BOM'

    safe_name = "".join(c for c in raw_name if c not in r'\/:*?"<>|').strip() or "Service_BOM"
    if not safe_name.lower().endswith('.xlsx'):
        filename = f"{safe_name}.xlsx"
    else:
        filename = safe_name

    out_path = os.path.join(EXPORT_DIR, filename)
    try:
        export_project_sbom(project_id, out_path)
        return send_file(out_path, as_attachment=True, download_name=filename)
    except Exception as e:
        return jsonify({'error': f'导出失败: {str(e)}'}), 500


@projects_bp.route('/api/project/<int:project_id>/export_offline_package', methods=['GET'])
def export_offline_package_api(project_id):
    try:
        zip_buf, zip_filename = generate_offline_package(project_id)
        return send_file(
            zip_buf,
            as_attachment=True,
            download_name=zip_filename,
            mimetype='application/zip'
        )
    except Exception as e:
        return jsonify({'success': False, 'error': f'生成现场离线运维资料包失败: {str(e)}'}), 500
