import os
import openpyxl
from .database import get_db
from .file_utils import sanitize_asset_name, safe_path_under, IMAGES_DIR

FIELD_DEFINITIONS = [
    {
        'key': 'material_no',
        'label': '物料号',
        'required': True,
        'exact_aliases': ['物料号', '物料编码', '料号', '物料代码', '零件号', '物料编号', '零件编号', '部件编号', '物料代号', 'material_no', 'material no', 'part no', 'item code'],
        'fuzzy_aliases': ['编码', '代码', 'code', 'part']
    },
    {
        'key': 'description',
        'label': '物料完整描述',
        'required': True,
        'exact_aliases': ['物料完整描述', '物料描述', '完整描述', '详细描述', '规格描述', '中文描述', '技术描述', '物料名称及规格', '名称及规格'],
        'fuzzy_aliases': ['描述', 'description', 'desc']
    },
    {
        'key': 'name',
        'label': '元器件名称',
        'required': False,
        'exact_aliases': ['元器件名称', '物料名称', '部件名称', '零件名称', '器件名称', '品名'],
        'fuzzy_aliases': ['名称', 'item name', 'name']
    },
    {
        'key': 'spec_model',
        'label': '规格型号',
        'required': False,
        'exact_aliases': ['规格型号', '型号/规格', '型号及规格', '型号规格', '规格/型号', '型号', '规格', '规格参数'],
        'fuzzy_aliases': ['spec_model', 'model', 'spec', 'specification']
    },
    {
        'key': 'manufacturer',
        'label': '制造商',
        'required': False,
        'exact_aliases': ['制造商', '生产厂家', '生产厂商', '生产厂家/品牌', '厂家', '生产商', '品牌', '供货商', '供应商', '制造厂商'],
        'fuzzy_aliases': ['manufacturer', 'mfr', 'vendor', 'brand']
    },
    {
        'key': 'unit',
        'label': '计量单位',
        'required': False,
        'exact_aliases': ['单位', '计量单位', '基本单位'],
        'fuzzy_aliases': ['unit', 'uom']
    },
    {
        'key': 'drawing_no',
        'label': '图纸/文件代号',
        'required': False,
        'exact_aliases': ['图纸/文件代号', '图纸代号', '图纸/代号', '图纸／代号', '文件代号', '图纸编号', '图号', '代号', '图纸/文件号', '图样代号'],
        'fuzzy_aliases': ['图纸', 'drawing_no', 'drawing no', 'drawing', '代号', '图号']
    },
    {
        'key': 'substitute',
        'label': '替代物料',
        'required': False,
        'exact_aliases': ['替代物料号/型号', '替代物料', '替代料', '替代物料号', '替代件', '替代料号'],
        'fuzzy_aliases': ['替代码', 'substitute', 'substitute_no', 'alt part']
    },
    {
        'key': 'notes',
        'label': '补充说明',
        'required': False,
        'exact_aliases': ['补充说明', '技术备注', '备注', '技术说明', '说明'],
        'fuzzy_aliases': ['notes', 'note', 'remark', 'comments']
    },
    {
        'key': 'level',
        'label': 'BOM层级 (用于总成结构)',
        'required': False,
        'exact_aliases': ['层级', 'BOM层级', '级别', '阶层'],
        'fuzzy_aliases': ['level', 'bom_level']
    },
    {
        'key': 'qty',
        'label': '单机数量',
        'required': False,
        'exact_aliases': ['单机数量', '数量', '装配数量', '单台数量', '用量'],
        'fuzzy_aliases': ['qty', 'quantity']
    }
]

def analyze_master_excel(file_path, target_sheet_name=None):
    """
    Analyzes an uploaded Excel workbook for Master Material importing.
    Identifies sheets, header row, and maps columns to standard field definitions.
    Determines certainty and provides preview data.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_sheets = wb.sheetnames

    # Filter out non-BOM sheets if multiple sheets exist
    candidate_sheets = [s for s in all_sheets if '封面' not in s and '签字' not in s and '修改' not in s]
    if not candidate_sheets:
        candidate_sheets = all_sheets

    # Determine sheet to analyze
    sheet_name = target_sheet_name if target_sheet_name in all_sheets else candidate_sheets[0]
    ws = wb[sheet_name]

    # Detect header row by scanning rows 1 to 8
    best_row = 1
    best_score = 0
    row_candidates = []

    for r in range(1, min(9, ws.max_row + 1)):
        row_score = 0
        non_empty = 0
        for c in range(1, min(30, ws.max_column + 1)):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip():
                non_empty += 1
                text = str(val).strip().lower()
                for f in FIELD_DEFINITIONS:
                    if any(alias in text for alias in f['exact_aliases']):
                        row_score += 2
                        break
                    elif any(alias in text for alias in f['fuzzy_aliases']):
                        row_score += 1
                        break
        row_candidates.append((r, row_score, non_empty))
        if row_score > best_score:
            best_score = row_score
            best_row = r

    # If no keywords matched, default to the first non-empty row or row 1
    if best_score == 0:
        for r, score, non_empty in row_candidates:
            if non_empty >= 2:
                best_row = r
                break

    header_row = best_row

    # Read headers
    headers = []
    for c in range(1, ws.max_column + 1):
        val = ws.cell(header_row, c).value
        title = str(val).strip() if val is not None else f"第{c}列"
        headers.append({'col_idx': c, 'title': title})

    # Column matching
    mapping = {}
    matched_scores = {}
    uncertain_reasons = []

    for f in FIELD_DEFINITIONS:
        f_key = f['key']
        best_col = None
        best_c_score = 0

        for h in headers:
            col_idx = h['col_idx']
            title = h['title'].strip().lower()
            if not title or title.startswith('第') and title.endswith('列'):
                continue

            # Exact alias match
            if any(title == alias.lower() for alias in f['exact_aliases']):
                score = 1.0
            elif any(alias.lower() in title for alias in f['exact_aliases']):
                score = 0.95
            elif any(title == alias.lower() for alias in f['fuzzy_aliases']):
                score = 0.85
            elif any(alias.lower() in title for alias in f['fuzzy_aliases']):
                score = 0.70
            else:
                score = 0.0

            if score > best_c_score:
                best_c_score = score
                best_col = col_idx

        if best_col and best_c_score >= 0.70:
            # Check for conflict with already matched field
            prev_field = next((k for k, v in mapping.items() if v == best_col), None)
            if prev_field:
                if best_c_score > matched_scores.get(prev_field, 0):
                    del mapping[prev_field]
                    mapping[f_key] = best_col
                    matched_scores[f_key] = best_c_score
            else:
                mapping[f_key] = best_col
                matched_scores[f_key] = best_c_score

    # Determine Certainty
    is_certain = True
    if 'material_no' not in mapping:
        is_certain = False
        uncertain_reasons.append('未在表格中自动匹配到"物料号"列，需人工指定')
    elif matched_scores.get('material_no', 0) < 0.95:
        is_certain = False
        uncertain_reasons.append(f'物料号匹配为 [{headers[mapping["material_no"]-1]["title"]}]，匹配度较低需确认')

    if 'description' not in mapping and 'name' not in mapping:
        is_certain = False
        uncertain_reasons.append('未自动识别出"物料完整描述"或"元器件名称"列')
    elif 'description' not in mapping:
        # Has name but no description
        is_certain = False
        uncertain_reasons.append('缺少"物料描述"列（已匹配名称），建议核对')

    if len(mapping) < 3:
        is_certain = False
        uncertain_reasons.append(f'仅自动匹配到 {len(mapping)} 个字段，需人工核对列对应关系')

    # Read preview rows (first 5 data rows)
    preview_rows = []
    total_valid_rows = 0
    mat_col = mapping.get('material_no')

    for r in range(header_row + 1, min(ws.max_row + 1, header_row + 2000)):
        row_vals = {}
        has_content = False
        for h in headers:
            val = ws.cell(r, h['col_idx']).value
            val_str = str(val).strip() if val is not None else ''
            row_vals[h['col_idx']] = val_str
            if val_str:
                has_content = True

        if not has_content:
            continue

        mat_val = row_vals.get(mat_col, '') if mat_col else ''
        # Exclude headers repeated in data or blank
        if mat_val and mat_val in ['物料号', '料号', 'None']:
            continue

        total_valid_rows += 1
        if len(preview_rows) < 5:
            preview_rows.append({
                'row_idx': r,
                'cells': row_vals
            })

    return {
        'success': True,
        'sheets': all_sheets,
        'current_sheet': sheet_name,
        'header_row': header_row,
        'headers': headers,
        'mapping': mapping,
        'is_certain': is_certain,
        'uncertain_reasons': uncertain_reasons,
        'preview_rows': preview_rows,
        'total_rows': total_valid_rows,
        'field_definitions': [{
            'key': f['key'],
            'label': f['label'],
            'required': f['required']
        } for f in FIELD_DEFINITIONS]
    }

def execute_master_import(file_path, sheet_name=None, header_row=None, mapping=None, overwrite=True):
    """
    Executes the Excel import into the Master Material Library using the confirmed column mapping.
    Also extracts hierarchy relationships and embedded images when present.
    """
    if not os.path.exists(file_path):
        return {'error': '上传的临时文件已失效，请重新选择文件上传'}

    wb = openpyxl.load_workbook(file_path, data_only=False)
    if not sheet_name or sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    if not header_row:
        header_row = 1
    else:
        header_row = int(header_row)

    if not mapping or not isinstance(mapping, dict):
        # Auto analyze if no mapping provided
        res = analyze_master_excel(file_path, sheet_name)
        mapping = res['mapping']

    mat_col = mapping.get('material_no')
    if not mat_col:
        return {'error': '必须映射"物料号"字段才能导入'}

    mat_col = int(mat_col)
    desc_col = int(mapping.get('description')) if mapping.get('description') else None
    name_col = int(mapping.get('name')) if mapping.get('name') else None
    spec_col = int(mapping.get('spec_model')) if mapping.get('spec_model') else None
    mfr_col = int(mapping.get('manufacturer')) if mapping.get('manufacturer') else None
    unit_col = int(mapping.get('unit')) if mapping.get('unit') else None
    dwg_col = int(mapping.get('drawing_no')) if mapping.get('drawing_no') else None
    sub_col = int(mapping.get('substitute')) if mapping.get('substitute') else None
    notes_col = int(mapping.get('notes')) if mapping.get('notes') else None
    lvl_col = int(mapping.get('level')) if mapping.get('level') else None
    qty_col = int(mapping.get('qty')) if mapping.get('qty') else None

    # Map embedded images by row number
    row_to_image = {}
    ws_images = getattr(ws, '_images', [])
    for img in ws_images:
        anchor_from = getattr(img, 'anchor', None)
        if anchor_from and hasattr(anchor_from, '_from'):
            r_idx = getattr(anchor_from._from, 'row', None)
            if r_idx is not None:
                row_to_image[r_idx + 1] = img

    conn = get_db()
    cursor = conn.cursor()

    count_inserted = 0
    count_updated = 0
    count_structures = 0

    hierarchy_stack = [] # [(level, parent_material_no)]

    for r in range(header_row + 1, ws.max_row + 1):
        val_mat = ws.cell(r, mat_col).value
        if val_mat is None:
            continue
        mat_no = str(val_mat).strip()
        if not mat_no or mat_no in ['物料号', '料号', 'None', 'nan']:
            continue

        desc = str(ws.cell(r, desc_col).value or '').strip() if desc_col else ''
        name = str(ws.cell(r, name_col).value or '').strip() if name_col else ''
        if not desc and name:
            desc = name
        elif not name and desc:
            name = desc

        spec_model = str(ws.cell(r, spec_col).value or '').strip() if spec_col else ''
        manufacturer = str(ws.cell(r, mfr_col).value or '').strip() if mfr_col else ''
        unit = str(ws.cell(r, unit_col).value or '').strip() if unit_col else 'PC'
        if not unit or unit.upper() == 'PCS': unit = 'PC'
        drawing_no = str(ws.cell(r, dwg_col).value or '').strip() if dwg_col else '/'
        if not drawing_no: drawing_no = '/'
        substitute = str(ws.cell(r, sub_col).value or '').strip() if sub_col else ''
        notes = str(ws.cell(r, notes_col).value or '').strip() if notes_col else ''
        qty = str(ws.cell(r, qty_col).value or '1').strip() if qty_col else '1'
        if not qty: qty = '1'

        # Extract image if present
        saved_img_rel = ''
        if r in row_to_image:
            try:
                img_obj = row_to_image[r]
                img_data = img_obj._data()
                safe_mat = sanitize_asset_name(mat_no, fallback='unnamed_img')
                img_filename = f"{safe_mat}.png"
                img_dest = safe_path_under(IMAGES_DIR, img_filename)
                if img_dest:
                    with open(img_dest, 'wb') as f_out:
                        f_out.write(img_data)
                    saved_img_rel = f"assets/images/{img_filename}"
            except Exception as e:
                print(f"Error saving image for {mat_no}: {e}")

        # Hierarchy handling if level column is mapped
        if lvl_col:
            raw_lvl = str(ws.cell(r, lvl_col).value or '').strip()
            lvl_num = 2
            lvl_str = '••'
            if '•••' in raw_lvl or '●●●' in raw_lvl or raw_lvl.startswith('3'):
                lvl_num = 3
                lvl_str = '•••'
            elif '••' in raw_lvl or '●●' in raw_lvl or raw_lvl.startswith('2'):
                lvl_num = 2
                lvl_str = '••'
            elif '•' in raw_lvl or '●' in raw_lvl or raw_lvl.startswith('1'):
                lvl_num = 1
                lvl_str = '•'

            while hierarchy_stack and hierarchy_stack[-1][0] >= lvl_num:
                hierarchy_stack.pop()

            if hierarchy_stack:
                parent_mat = hierarchy_stack[-1][1]
                if parent_mat and parent_mat != mat_no:
                    # Check if structure link already exists
                    cursor.execute('''
                    SELECT id FROM material_structures
                    WHERE parent_material_no = ? AND child_material_no = ?
                    ''', (parent_mat, mat_no))
                    existing_struct = cursor.fetchone()
                    if not existing_struct:
                        cursor.execute('SELECT COALESCE(MAX(sort_order), 0) + 1 FROM material_structures WHERE parent_material_no = ?', (parent_mat,))
                        next_sort = cursor.fetchone()[0]
                        cursor.execute('''
                        INSERT INTO material_structures (parent_material_no, child_material_no, child_level_num, child_level_str, sort_order, qty, unit, drawing_no, notes, substitute)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        ''', (parent_mat, mat_no, lvl_num, lvl_str, next_sort, qty, unit, drawing_no, notes, substitute))
                        count_structures += 1
            hierarchy_stack.append((lvl_num, mat_no))

        # Check existing material
        cursor.execute('SELECT material_no, image_path, manual_path FROM materials WHERE material_no = ?', (mat_no,))
        row = cursor.fetchone()

        if row:
            if overwrite:
                final_img = saved_img_rel if saved_img_rel else row['image_path']
                cursor.execute('''
                UPDATE materials
                SET name = CASE WHEN ? != '' THEN ? ELSE name END,
                    description = CASE WHEN ? != '' THEN ? ELSE description END,
                    spec_model = CASE WHEN ? != '' THEN ? ELSE spec_model END,
                    manufacturer = CASE WHEN ? != '' THEN ? ELSE manufacturer END,
                    unit = CASE WHEN ? != '' THEN ? ELSE unit END,
                    drawing_no = CASE WHEN ? != '' AND ? != '/' THEN ? ELSE drawing_no END,
                    substitute = CASE WHEN ? != '' THEN ? ELSE substitute END,
                    notes = CASE WHEN ? != '' THEN ? ELSE notes END,
                    image_path = ?
                WHERE material_no = ?
                ''', (name, name, desc, desc, spec_model, spec_model, manufacturer, manufacturer,
                      unit, unit, drawing_no, drawing_no, drawing_no, substitute, substitute, notes, notes, final_img, mat_no))
                count_updated += 1
        else:
            cursor.execute('''
            INSERT INTO materials (material_no, name, description, spec_model, manufacturer, unit, drawing_no, substitute, notes, image_path)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (mat_no, name, desc, spec_model, manufacturer, unit, drawing_no, substitute, notes, saved_img_rel))
            count_inserted += 1

    conn.commit()
    conn.close()

    total_processed = count_inserted + count_updated
    struct_msg = f"，自动识别并构建了 {count_structures} 条总成层级关系" if count_structures > 0 else ""
    return {
        'success': True,
        'imported_count': count_inserted,
        'updated_count': count_updated,
        'total_processed': total_processed,
        'structures_created': count_structures,
        'message': f'成功从 Excel 导入 {total_processed} 项物料（新增 {count_inserted} 项，更新 {count_updated} 项{struct_msg}）！'
    }
