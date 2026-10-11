import os
import io
import zipfile
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from .database import get_db, find_template_path

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def ensure_company_logo(tpl_path):
    """
    Extracts the company logo from the template file if not already present.
    """
    logo_path = os.path.join(BASE_DIR, 'assets', 'company_logo.png')
    if not os.path.exists(logo_path) and tpl_path and os.path.exists(tpl_path):
        try:
            with zipfile.ZipFile(tpl_path, 'r') as z:
                # In template, xl/media/image1.png is the company logo
                if 'xl/media/image1.png' in z.namelist():
                    os.makedirs(os.path.dirname(logo_path), exist_ok=True)
                    with open(logo_path, 'wb') as f:
                        f.write(z.read('xl/media/image1.png'))
        except Exception as e:
            print("Warning extracting logo:", e)
    return logo_path

def export_project_sbom(project_id, output_path, embed_manuals=False):
    """
    Exports the project and all its subsystems to a standardized Excel workbook
    strictly preserving the original template's Cover (封面), Signatures & ECN (签字及修改页),
    formulas, merged cells, fonts, borders, layouts, and company logos.
    """
    conn = get_db()
    cursor = conn.cursor()

    # 1. Fetch project info
    cursor.execute('SELECT * FROM projects WHERE id = ?', (project_id,))
    project = cursor.fetchone()
    if not project:
        conn.close()
        raise ValueError(f"未找到ID为 {project_id} 的项目。")

    # 2. Locate and load the original template
    tpl_path = find_template_path()
    if not tpl_path or not os.path.exists(tpl_path):
        conn.close()
        raise FileNotFoundError("未在'Excel 模版'目录或根目录下找到SBOM模版文件，请确认模版文件存在！")

    wb = openpyxl.load_workbook(tpl_path)
    company_logo_path = ensure_company_logo(tpl_path)

    # Standard borders and fonts for data cells
    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )
    song_font = Font(name='宋体', size=10)
    tnr_font = Font(name='Times New Roman', size=10)
    level_font = Font(name='Times New Roman', size=11)
    link_font = Font(name='微软雅黑', size=9, color='0055CC', underline='single')

    # ================= 1. 封面 (Cover) =================
    # ================= 1. 封面 (Cover) =================
    cover_sheet_name = '封面' if '封面' in wb.sheetnames else ([s for s in wb.sheetnames if '封面' in s] or ['封面'])[0]
    ws_cover = wb[cover_sheet_name]

    # V2 模版封面布局：
    # A2: 标题/部套大类（如：风力发电机组）
    ws_cover['A2'].value = project['name'] or '风力发电机组'
    # A3: 平台（举例：WL，若无平台则留空）
    ws_cover['A3'].value = (project['platform'] if 'platform' in project.keys() and project['platform'] else '') or ''
    # E3: 机型（举例：EW6.25-220。注释说明：平台通用的话，机型不用填）
    ws_cover['E3'].value = project['code'] or 'EW6.25-220'
    # A4: 供应商名称+部套名称+SBOM+型号（举例：变桨轴承润滑系统SBOM_40116237600。注释：属于系统集成，没有型号的，型号可以省略）
    ws_cover['A4'].value = project['system_name'] or f"{project['code']}变桨轴承润滑系统SBOM"
    # A5: 文件号 (DFD...)
    ws_cover['A5'].value = project['doc_no'] or ''
    # E5: ECN 号
    ws_cover['E5'].value = project['ecn_no'] or ''
    # A7: 公司名称
    ws_cover['A7'].value = project['company'] or '上海电气风电集团股份有限公司'
    # Row 8: 年 / 月 (C8=年, E8=月)
    now = datetime.now()
    ws_cover['C8'].value = now.year
    ws_cover['E8'].value = now.month
    # A11: 表单编号
    ws_cover['A11'].value = project['form_no'] or '  FDJL-JS-037'

    # ================= 2. 签字页 (Signatures) =================
    if '签字页' in wb.sheetnames:
        ws_sig = wb['签字页']
        # Note: A3 包含公式 =封面!A4（注释：直接引用封面文件名），J1 包含公式 =封面!A5，保持不变
        cursor.execute('SELECT * FROM signatures WHERE project_id = ?', (project_id,))
        sigs = {s['role']: (s['sign_name'], s['sign_date']) for s in cursor.fetchall()}

        # 模版行号：编制(11), 校对/校核(14), 审核(17), 会签(20), 标审(23), 审定(26)
        role_map = [
            (['编制'], 11),
            (['校对', '校核'], 14),
            (['审核'], 17),
            (['会签'], 20),
            (['标审'], 23),
            (['审定'], 26),
        ]
        for role_aliases, r in role_map:
            name_val, date_val = '', ''
            for r_alias in role_aliases:
                if r_alias in sigs:
                    name_val, date_val = sigs[r_alias]
                    break
            ws_sig[f'F{r}'].value = name_val or None
            ws_sig[f'I{r}'].value = date_val or None

    # ================= 3. 更改页 (ECN Records) =================
    if '更改页' in wb.sheetnames:
        ws_change = wb['更改页']
        # Note: A3 包含公式 =封面!A4（注释：直接引用封面文件名），L1 包含公式 =封面!A5，保持不变
        # 清空模版占位行 (第8至32行)
        for r in range(8, 33):
            ws_change[f'B{r}'].value = None
            ws_change[f'F{r}'].value = None
            ws_change[f'H{r}'].value = None
            ws_change[f'K{r}'].value = None
            ws_change[f'M{r}'].value = None

        cursor.execute('SELECT * FROM ecn_records WHERE project_id = ? ORDER BY id', (project_id,))
        ecn_list = cursor.fetchall()
        for i, ecn in enumerate(ecn_list):
            r = 8 + i
            if r <= 32:
                ws_change[f'B{r}'].value = ecn['ecn_no']
                ws_change[f'F{r}'].value = ecn['change_pos'] if 'change_pos' in ecn.keys() else ''
                ws_change[f'H{r}'].value = ecn['change_desc']
                ws_change[f'K{r}'].value = ecn['changer']
                ws_change[f'M{r}'].value = ecn['change_date']

    # ================= 4. BOM Subsystem Sheets =================
    cursor.execute('SELECT * FROM subsystems WHERE project_id = ? ORDER BY sort_order, id', (project_id,))
    subsystems = cursor.fetchall()

    # Identify prototype BOM template sheet (e.g. 盘古SBOM or first BOM sheet)
    proto_candidates = [name for name in wb.sheetnames if name not in ['封面', '签字页', '更改页']]
    proto_sheet = wb[proto_candidates[0]] if proto_candidates else None

    # Track sheet names in this export
    exported_sub_sheets = []

    for sub in subsystems:
        sub_name = sub['name']
        exported_sub_sheets.append(sub_name)

        if sub_name in wb.sheetnames:
            ws = wb[sub_name]
            # Keep company logo (Image 0), clear old template product sample images
            if len(ws._images) > 0:
                ws._images = [ws._images[0]]
            else:
                ws._images = []
        elif proto_sheet is not None:
            # Create sheet by copying prototype template sheet
            ws = wb.copy_worksheet(proto_sheet)
            ws.title = sub_name
            # Copy company logo to new sheet
            if os.path.exists(company_logo_path):
                logo_img = OpenpyxlImage(company_logo_path)
                logo_img.width = 212
                logo_img.height = 76
                ws.add_image(logo_img, 'A1')
        else:
            ws = wb.create_sheet(title=sub_name)

        # Clear old template sample data rows 4+
        if ws.max_row >= 4:
            ws.delete_rows(4, ws.max_row - 3)

        # Update headers in row 1 & row 2
        ws['F1'].value = project['company'] or '上海电气风电集团股份有限公司'
        sub_doc_no = (sub['doc_no'] or '').strip() or (project['doc_no'] or '').strip() or '/'
        ws['I1'].value = sub_doc_no
        ws['M1'].value = project['ecn_no'] or ''
        
        if sub['system_title']:
            ws['I2'].value = sub['system_title']
        elif sub_name == 'SBOM':
            ws['I2'].value = project['system_name'] or f"{project['code']}变桨轴承润滑系统"
        elif sub_name.endswith('SBOM'):
            ws['I2'].value = f"{project['code']}{sub_name[:-4]}变桨轴承润滑系统"
        else:
            ws['I2'].value = f"{project['code']}{sub_name}润滑系统"

        # Standardize Row 3 header fonts: ensure uniform 12pt bold font across all columns
        for c in range(1, 17):
            h_cell = ws.cell(row=3, column=c)
            if h_cell.value:
                cur_name = h_cell.font.name if (h_cell.font and h_cell.font.name in ['Times New Roman', '宋体']) else 'Times New Roman'
                h_cell.font = Font(
                    name=cur_name,
                    size=12.0,
                    bold=True,
                    color=h_cell.font.color if (h_cell.font and h_cell.font.color) else None
                )

        # Fetch BOM items with material master details
        cursor.execute('''
        SELECT b.*, m.name as m_name, m.description as m_desc, m.spec_model as m_model,
               m.manufacturer as m_mfr, m.unit as m_unit, m.image_path, m.manual_path,
               m.drawing_no as m_drawing_no
        FROM bom_items b
        LEFT JOIN materials m ON b.material_no = m.material_no
        WHERE b.subsystem_id = ?
        ORDER BY b.row_num, b.id
        ''', (sub['id'],))
        bom_items = cursor.fetchall()

        # Check duplicate material numbers within this subsystem
        sub_mat_counts = {}
        for it in bom_items:
            m_no = (it['material_no'] or '').strip()
            if m_no:
                sub_mat_counts[m_no] = sub_mat_counts.get(m_no, 0) + 1

        dup_mat_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        dup_mat_font = Font(name='Times New Roman', size=10, bold=True, color='9A3412')

        # Dynamically detect image column from Row 3 headers
        from openpyxl.utils import get_column_letter
        img_col_idx = 13  # Default Column N (0-based 13)
        img_col_letter = 'N'
        for c in range(1, 26):
            header_val = str(ws.cell(row=3, column=c).value or '').strip()
            if any(k in header_val for k in ['图片', '照片', '实物图', '外观图']) or header_val == '图':
                img_col_idx = c - 1
                img_col_letter = get_column_letter(c)
                break

        # Calculate actual cell width in pixels from template column dimensions
        col_dim = ws.column_dimensions.get(img_col_letter)
        col_w_unit = float(col_dim.width) if (col_dim and col_dim.width is not None) else 29.0
        # Standard Excel 96 DPI: 1 width unit ≈ 7.5px + 5px padding
        cell_w = max(40, int(col_w_unit * 7.5 + 5))

        for idx, item in enumerate(bom_items):
            r = 4 + idx
            ws.row_dimensions[r].height = 70  # Standard ample height for clear images
            row_h_pt = float(ws.row_dimensions[r].height or 70)
            cell_h = max(30, int(row_h_pt * 96 / 72))

            level_num_val = item['level_num'] or 1
            level_str_val = item['level_str'] or ('•' * level_num_val)
            is_dup_row = sub_mat_counts.get((item['material_no'] or '').strip(), 0) > 1

            row_data = [
                ('A', item['row_num'], tnr_font, Alignment(horizontal='center', vertical='center')),
                ('B', f'=REPT("•",C{r})', level_font, Alignment(horizontal='right', vertical='center')),
                ('C', level_num_val, level_font, Alignment(horizontal='left', vertical='center')),
                ('D', item['seq_num'] or (item['row_num'] * 10), tnr_font, Alignment(horizontal='center', vertical='center')),
                ('E', item['material_no'], tnr_font, Alignment(horizontal='center', vertical='center')),
                ('F', item['m_desc'] or '', song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('G', item['m_name'] or '', song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('H', (item['drawing_no'] if (item['drawing_no'] and item['drawing_no'] != '/') else (item['m_drawing_no'] or '/')), song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('I', item['m_model'] or '', song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('J', item['m_mfr'] or '', song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('K', item['qty'], song_font, Alignment(horizontal='center', vertical='center')),
                ('L', item['unit'] or item['m_unit'] or 'PC', tnr_font, Alignment(horizontal='center', vertical='center')),
                ('M', '', tnr_font, Alignment(horizontal='center', vertical='center')),
                ('N', '', tnr_font, Alignment(horizontal='center', vertical='center')),
                ('O', item['notes'] or '', song_font, Alignment(horizontal='center', vertical='center', wrap_text=True)),
                ('P', item['substitute'] or '', song_font, Alignment(horizontal='left', vertical='center', wrap_text=True))
            ]

            for col_letter, val, font_obj, align_obj in row_data:
                cell = ws[f'{col_letter}{r}']
                cell.value = val
                cell.font = font_obj
                cell.alignment = align_obj
                cell.border = thin_border
                if is_dup_row and col_letter == 'E':
                    cell.fill = dup_mat_fill
                    cell.font = dup_mat_font

            # Column M: Manual Link / Info (Hyperlink mode)
            manual_path = item['manual_path']
            if manual_path and not embed_manuals:
                full_m_path = os.path.join(BASE_DIR, manual_path)
                if os.path.exists(full_m_path):
                    m_cell = ws[f'M{r}']
                    m_cell.value = '查看技术手册'
                    m_cell.font = link_font
                    m_cell.alignment = Alignment(horizontal='center', vertical='center')
                    # Relative hyperlink for clean opening in Excel
                    m_cell.hyperlink = manual_path.replace('/', '\\')

            # Image Embedding (Proportionally resized and dynamically centered into cell)
            img_path = item['image_path']
            if img_path:
                full_img_path = os.path.join(BASE_DIR, img_path)
                if os.path.exists(full_img_path) and os.path.splitext(full_img_path)[1].lower() in ['.png', '.jpg', '.jpeg']:
                    try:
                        with PILImage.open(full_img_path) as pil_img:
                            orig_w, orig_h = pil_img.size
                            # Safe margins: leave 8px horizontally on each side, 6px vertically
                            max_w = max(20, cell_w - 16)
                            max_h = max(20, cell_h - 12)
                            scale = min(max_w / orig_w, max_h / orig_h)
                            new_w = max(1, int(orig_w * scale))
                            new_h = max(1, int(orig_h * scale))

                            # 优化嵌入体积：若原图分辨率过大，进行 Retina 双倍清晰度优化压缩（彻底解决 26MB+ 巨型 Excel 膨胀与内存爆满）
                            retina_w = max(new_w * 2, 260)
                            retina_h = max(new_h * 2, 180)
                            file_size = os.path.getsize(full_img_path)
                            if (orig_w > retina_w or orig_h > retina_h or file_size > 150 * 1024):
                                thumb_img = pil_img.copy()
                                thumb_img.thumbnail((retina_w, retina_h), PILImage.Resampling.LANCZOS)
                                if thumb_img.mode in ('RGBA', 'P', 'LA'):
                                    bg = PILImage.new('RGB', thumb_img.size, (255, 255, 255))
                                    if thumb_img.mode == 'P':
                                        thumb_img = thumb_img.convert('RGBA')
                                    bg.paste(thumb_img, mask=thumb_img.split()[-1] if 'A' in thumb_img.mode else None)
                                    thumb_img = bg
                                img_bio = io.BytesIO()
                                thumb_img.save(img_bio, format='JPEG', quality=84, optimize=True)
                                img_bio.seek(0)
                                excel_img = OpenpyxlImage(img_bio)
                            else:
                                excel_img = OpenpyxlImage(full_img_path)

                        excel_img.width = new_w
                        excel_img.height = new_h

                        # Calculate horizontal and vertical centering offset inside the actual cell dimensions
                        offset_x = max(0, (cell_w - new_w) // 2)
                        offset_y = max(0, (cell_h - new_h) // 2)

                        marker = AnchorMarker(
                            col=img_col_idx,
                            colOff=pixels_to_EMU(offset_x),
                            row=r - 1,
                            rowOff=pixels_to_EMU(offset_y)
                        )
                        size = XDRPositiveSize2D(pixels_to_EMU(new_w), pixels_to_EMU(new_h))
                        excel_img.anchor = OneCellAnchor(_from=marker, ext=size)
                        ws.add_image(excel_img)
                    except Exception as e:
                        print(f"Warning: failed to add image for {item['material_no']}: {e}")


    # ================= 4. Cleanup and Reorder Sheets (Strictly V2) =================
    # Remove any extra template sample BOM sheets that are not in the current project's subsystems
    standard_sheets = ['封面', '签字页', '更改页']
    keep_sheets = set(standard_sheets + exported_sub_sheets)
    sheets_to_remove = [name for name in wb.sheetnames if name not in keep_sheets]
    for sname in sheets_to_remove:
        wb.remove(wb[sname])

    # Ensure precise ordering: 封面 -> 签字页 -> 更改页 -> 各子系统
    desired_order = [s for s in standard_sheets if s in wb.sheetnames] + [s for s in exported_sub_sheets if s in wb.sheetnames]
    for s in wb.sheetnames:
        if s not in desired_order:
            desired_order.append(s)

    wb._sheets = [wb[name] for name in desired_order]

    conn.close()
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    print(f"Exported successfully to: {output_path}")

    if embed_manuals:
        embed_manuals_into_excel(output_path, project_id)

    return output_path


def get_pdf_icon():
    """
    获取系统中真实的 PDF 专用图标文件及图标索引。
    1. 优先使用本地内置的高清红标 PDF 专用图标 assets/pdf_icon.ico
    2. 其次查询 Windows 注册表当前默认关联的 PDF 应用程序官方文档图标
    3. 常见 Foxit / Adobe Acrobat 安装路径兜底
    """
    local_ico = os.path.join(BASE_DIR, 'assets', 'pdf_icon.ico')
    if os.path.exists(local_ico):
        return local_ico, 0

    try:
        import winreg
        for subkey in [r'Software\Microsoft\Windows\CurrentVersion\Explorer\FileExts\.pdf\UserChoice', r'.pdf']:
            try:
                root = winreg.HKEY_CURRENT_USER if 'UserChoice' in subkey else winreg.HKEY_CLASSES_ROOT
                with winreg.OpenKey(root, subkey) as k:
                    prog_id = winreg.QueryValueEx(k, 'ProgId' if 'UserChoice' in subkey else '')[0]
                    if prog_id:
                        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, rf'{prog_id}\DefaultIcon') as ik:
                            raw_val = winreg.QueryValueEx(ik, '')[0].strip().strip('"')
                            parts = raw_val.split(',')
                            path = parts[0].strip().strip('"')
                            idx = int(parts[1].strip()) if len(parts) > 1 else 0
                            if os.path.exists(path):
                                return path, idx
            except Exception:
                pass
    except Exception:
        pass

    candidates = [
        (r'C:\Program Files (x86)\Foxit Software\Foxit PDF Editor\FoxitPDFEditor.exe', 1),
        (r'C:\Program Files\Foxit Software\Foxit PDF Editor\FoxitPDFEditor.exe', 1),
        (r'C:\Program Files\Adobe\Acrobat DC\Acrobat\Acrobat.exe', 1),
        (r'C:\Program Files (x86)\Adobe\Acrobat Reader DC\Reader\AcroRd32.exe', 0),
    ]
    for p, idx in candidates:
        if os.path.exists(p):
            return p, idx

    return r'C:\Windows\System32\packager.dll', 0


def embed_manuals_into_excel(excel_path, project_id):
    """
    通过 Windows Excel COM 自动化将零部件关联的 PDF 技术手册以 OLE 对象直接嵌入到 Excel 表格 M 列中。
    嵌入后，Excel 成为自包含独立文件，接收者双击单元格内 PDF 图标即可直接打开查看手册。
    """
    abs_excel_path = os.path.abspath(excel_path)
    if not os.path.exists(abs_excel_path):
        raise FileNotFoundError(f"Excel 文件不存在: {abs_excel_path}")

    # 查询该项目中所有挂载了有效手册的子系统物料
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT s.name as sheet_name, b.row_num, m.material_no, m.name as mat_name, m.manual_path
        FROM bom_items b
        JOIN subsystems s ON b.subsystem_id = s.id
        LEFT JOIN materials m ON b.material_no = m.material_no
        WHERE s.project_id = ? AND m.manual_path IS NOT NULL AND m.manual_path != ''
        ORDER BY s.sort_order, s.id, b.row_num
    ''', (project_id,))
    manual_rows = cursor.fetchall()
    conn.close()

    if not manual_rows:
        return abs_excel_path

    import gc
    try:
        import comtypes.client
    except ImportError:
        raise RuntimeError("系统未安装 comtypes 库，无法执行 Excel OLE 对象嵌入。")

    icon_file, icon_index = get_pdf_icon()
    print(f"Using PDF icon: {icon_file}, index: {icon_index}")

    excel = None
    wb = None
    try:
        excel = comtypes.client.CreateObject('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False

        wb = excel.Workbooks.Open(abs_excel_path)
        embedded_count = 0

        for row in manual_rows:
            sheet_name = row['sheet_name']
            row_idx = 4 + (row['row_num'] - 1)
            m_path = row['manual_path']
            full_m_path = os.path.abspath(m_path)
            if not os.path.exists(full_m_path):
                continue

            try:
                ws = wb.Sheets(sheet_name)
            except Exception:
                continue

            cell = ws.Range(f'M{row_idx}')
            cell.ClearContents()

            lbl = os.path.basename(m_path)
            if len(lbl) > 22:
                lbl = lbl[:20] + '..'

            try:
                # OLEObjects.Add(ClassType, Filename, Link, DisplayAsIcon, IconFileName, IconIndex, IconLabel, Left, Top, Width, Height)
                ole = ws.OLEObjects().Add(
                    None,
                    full_m_path,
                    False,  # Link = False (真正嵌入为内部二进制数据包)
                    True,   # DisplayAsIcon = True
                    icon_file,
                    icon_index,
                    lbl,
                    float(cell.Left),
                    float(cell.Top),
                    60,
                    45
                )

                # 精确计算水平和垂直居中位置，放置在单元格正中间
                ole_w = float(ole.Width)
                ole_h = float(ole.Height)
                ole.Left = float(cell.Left) + max(0.0, (float(cell.Width) - ole_w) / 2.0)
                ole.Top = float(cell.Top) + max(0.0, (float(cell.Height) - ole_h) / 2.0)
                ole.Placement = 2  # xlMove: 单元格变动或增删行时位置随单元格移动且不失真变形
                embedded_count += 1
            except Exception as e:
                print(f"Warning: 嵌入手册 {lbl} 失败: {e}")

        # 使用 SaveAs 51 (xlOpenXMLWorkbook) 保存，确保 OLE 复合二进制流完整写入
        temp_saved_path = abs_excel_path + '.embedded.xlsx'
        if os.path.exists(temp_saved_path):
            try: os.remove(temp_saved_path)
            except: pass
        wb.SaveAs(temp_saved_path, 51)
        wb.Close(False)
        wb = None
        excel.Quit()
        excel = None
        gc.collect()

        # 替换原文件
        if os.path.exists(temp_saved_path):
            if os.path.exists(abs_excel_path):
                os.remove(abs_excel_path)
            os.rename(temp_saved_path, abs_excel_path)

        print(f"成功将 {embedded_count} 份技术手册嵌入至 Excel: {abs_excel_path}")
        return abs_excel_path
    except Exception as e:
        if wb:
            try: wb.Close(False)
            except: pass
        if excel:
            try: excel.Quit()
            except: pass
        gc.collect()
        raise RuntimeError(f"Excel 文档嵌入处理失败: {str(e)}")


if __name__ == '__main__':
    test_out = os.path.join(BASE_DIR, 'exports', 'test_export_sbom.xlsx')
    export_project_sbom(1, test_out)
